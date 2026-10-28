from __future__ import annotations

import datetime as dt
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from big_export import REQUIRED_HEADERS
from daily_raw04 import build_daily_raw04_tables
from io_excel import write_output
from main import (
    _build_daily_raw04_header_table,
    _build_raw04_primary_without_report,
    _merge_primary_and_raw04_tables,
)


class DailyRaw04Tests(unittest.TestCase):

    def test_header_uses_certificate_table_totals_in_prose_order(self) -> None:
        header = _build_daily_raw04_header_table([
            {"table": "Таблица1", "auto_rows": [{"РФ": 10, "РБ": 2, "КР": 1, "РА": 0}], "rail_rows": [{"РФ": 3, "РБ": 4, "КР": 0, "РА": 0}]},
            {"table": "Таблица2", "export_split_rows": [{"Страна": "Российская Федерация", "Итого": 5}, {"Страна": "Республика Беларусь", "Итого": 2}, {"Страна": "Кыргызская Республика", "Итого": 1}, {"Страна": "Республика Армения", "Итого": 0}, {"Итого": 8}], "trade_split_rows": [{"Страна": "Российская Федерация", "Итого": 7}, {"Страна": "Республика Беларусь", "Итого": 3}, {"Страна": "Кыргызская Республика", "Итого": 4}, {"Страна": "Республика Армения", "Итого": 0}, {"Итого": 14}]},
            {"table": "Таблица3", "rows": [{"Итого": 9}]},
            {"table": "Таблица4", "auto_in_rows": [{"Количество": 20}], "rail_in_rows": [{"Количество": 6}]},
        ])
        self.assertEqual(header["table"], "Шапка")
        self.assertEqual(header["rows"][0]["Число"], 42)
        self.assertEqual(header["rows"][1]["Число"], 9)
        self.assertEqual(header["rows"][2]["Число"], 26)
        self.assertEqual(header["rows"][3]["Число"], 20)
        self.assertEqual(header["rows"][4]["Число"], 6)
        self.assertEqual(header["rows"][12]["Число"], 4)

    def test_combined_mode_uses_primary_mode_tables_and_certificate_order(self) -> None:
        old = [
            {"table": "Таблица1", "transport_split": True},
            {"table": "Таблица2", "transport_split": True},
            {"table": "Таблица4"}, {"table": "Таблица5", "rows": []},
            {"table": "Таблица6", "rows": []},
        ]
        raw = [
            {"table": "Таблица3"}, {"table": "Таблица6"},
            {"table": "Таблица7"}, {"table": "Диагностика"},
        ]
        tables = _merge_primary_and_raw04_tables(old, raw)
        self.assertEqual([table["table"] for table in tables], [
            "Шапка", "Таблица1", "Таблица2", "Таблица3", "Таблица4", "Таблица5",
            "Таблица6", "Таблица7", "Таблица8", "Диагностика",
        ])
        self.assertTrue(tables[8]["legacy_stock_layout"])
        self.assertTrue(tables[1]["transport_split"])
        self.assertTrue(tables[2]["transport_split"])
    def _source(self, path: Path) -> None:
        headers = [
            "Номер перевозки", "Дата обновления", "Статус", "Время статуса",
            "Дата создания", "Дата активации", "Страна начала", "Страна завершения",
            "Промежуточные страны", "Начальный пункт", "Тип перевозки",
            "Тип транспорта", "Номера НП",
        ]
        rows = [
            ["TR-1", "2026-10-01 10:00:00", "Активированная перевозка", "2026-10-01 10:00:00", "2026-09-20", "2026-09-21", "РК", "РФ", "", "ТП Алаколь", "Таможенная процедура таможенного транзита", "Автомобильный транспорт", "0001; 0001"],
            ["EX-1", "2026-10-01 10:00:00", "Активированная перевозка", "2026-10-01 10:00:00", "2026-09-20", "2026-09-21", "РК", "РБ", "", "Темиржол", "Экспорт", "Железнодорожный траспорт", "0002"],
            ["MT-1", "2026-10-01 10:00:00", "Завершенная перевозка", "2026-10-01 10:00:00", "2026-09-20", "", "РК", "РК", "", "Астана", "Взаимной торговли", "Автомобильный транспорт", "0003"],
            ["IN-1", "2026-10-01 10:00:00", "Завершенная перевозка", "2026-10-01 11:00:00", "2026-09-20", "2026-09-21", "РФ", "РК", "", "", "Транзит", "Железнодорожный транспорт", "0004"],
            ["UK-1", "2026-10-01 10:00:00", "Активированная перевозка", "2026-10-01 10:00:00", "2026-09-20", "2026-09-21", "РК", "РФ", "", "Неизвестный пост", "Транзит", "Автомобильный транспорт", "0005"],
            ["WAIT-1", "2026-10-01 10:00:00", "Ожидание навешивания", "2026-10-01 10:00:00", "2026-09-20", "2026-09-21", "РК", "РФ", "", "ТП Алаколь", "Транзит", "Автомобильный транспорт", "0006"],
            ["TRACK-ACTIVE", "2026-10-01 10:00:00", "Активированная перевозка", "2026-10-01 10:00:00", "2026-02-11", "2026-02-11", "РБ", "КР", "РК", "", "Транзит", "Автомобильный транспорт", ""],
            ["TRACK-FINISHED", "2026-10-01 10:00:00", "Завершенная перевозка", "2026-10-01 10:00:00", "2026-02-11", "2026-02-11", "РБ", "КР", "РК", "", "Транзит", "Автомобильный транспорт", ""],
        ]
        book = Workbook()
        ws = book.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        book.save(path)

    def test_builds_certificate_tables_and_reconciles_seals(self) -> None:
        registry = Path(__file__).parents[1] / "src" / "data" / "border_crossing_points.json"
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "raw04.xlsx"
            output = Path(temp_dir) / "result.xlsx"
            self._source(source)
            tables = build_daily_raw04_tables(source, registry)
            self.assertEqual([table["table"] for table in tables], [
                "Таблица1", "Таблица2", "Таблица3", "Таблица4", "Таблица5", "Таблица6", "Таблица7", "Диагностика",
            ])
            self.assertEqual(tables[0]["auto_rows"][0]["РФ"], 1)
            self.assertEqual(tables[1]["export_split_rows"][1]["ЖД"], 1)
            self.assertEqual(tables[2]["rows"][0]["РБ - КР"], 2)
            self.assertEqual(tables[3]["rail_in_rows"][0]["Количество"], 1)
            self.assertEqual(tables[5]["rows"][-2]["Итого"], 591)
            self.assertEqual(tables[5]["rows"][-2]["Транзит"], 35)
            self.assertEqual(tables[5]["rows"][-2]["Экспорт"], 19)
            self.assertEqual(tables[5]["rows"][-2]["Взаимная торговля"], 537)
            self.assertEqual(tables[5]["rows"][-1]["Итого"], 594)
            self.assertEqual(tables[6]["total"], 2)
            self.assertEqual(tables[6]["unresolved_points"], 1)
            write_output(output, tables)
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(workbook.sheetnames, [
                "Таблица1", "Таблица2", "Таблица3", "Таблица4", "Таблица5", "Таблица6", "Таблица7", "Диагностика",
            ])
            workbook.close()

    def test_raw04_runs_without_optional_report_and_marks_stock_unavailable(self) -> None:
        """Only the large export is necessary for the primary RAW04 mode."""
        registry = Path(__file__).parents[1] / "src" / "data" / "border_crossing_points.json"
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "raw04.xlsx"
            output = Path(temp_dir) / "result.xlsx"
            headers = list(REQUIRED_HEADERS)
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(headers)
            record = {header: "" for header in headers}
            record.update({
                "Дата создания": "2026-10-01 10:00:00",
                "Номер перевозки": "ONLY-BIG-1",
                "Статус": "Завершенная перевозка",
                "Дата обновления": "2026-10-01 10:00:00",
                "Тип перевозки": "Таможенная процедура таможенного транзита",
                "Тип транспорта": "Автомобильный транспорт",
                "Владелец": "КЕДЕН KEDEN",
                "Компания": "Тест",
                "БИН": "123",
                "Номера ТС": "TEST",
                "Номера НП": "0001",
                "Кол. путевых точек": "1",
                "Кол. водителей": "1",
                "Кол. операторов": "1",
                "Начальный пункт": "ТП Алаколь",
                "Конечный пункт": "Россия",
                "Страна начала": "Республика Казахстан",
                "Страна завершения": "Российская Федерация",
            })
            worksheet.append([record[header] for header in headers])
            workbook.save(source)

            primary, code = _build_raw04_primary_without_report(source)
            self.assertEqual(code, 0)
            self.assertIsNotNone(primary)
            self.assertTrue(primary[4]["stock_unavailable"])

            raw = build_daily_raw04_tables(source, registry)
            tables = _merge_primary_and_raw04_tables(primary, raw)
            write_output(output, tables)

            workbook = load_workbook(output, read_only=True)
            self.assertEqual(workbook.sheetnames[0], "Шапка")
            self.assertIn("Таблица8", workbook.sheetnames)
            self.assertIn("Нет файла report_*.xlsx", workbook["Таблица8"]["A2"].value)
            workbook.close()

