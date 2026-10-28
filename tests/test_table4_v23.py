from __future__ import annotations

import datetime as dt
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"
sys.path.insert(0, str(SRC_DIR))

from io_excel import PILOT_KARASU_NP, _stage2_elapsed_days, _write_table4  # noqa: E402
from builders import build_table4  # noqa: E402
from config import load_config  # noqa: E402


class Table4Version23Tests(unittest.TestCase):
    def test_left_stock_table_uses_selected_source_rows_without_duplicate_zko(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "СЕНТЯБРЬ"
        worksheet.append(["Пост", "01.09.2026"])
        worksheet.append(["Алматинская область", 1454])
        worksheet.append(["Область Улытау", 0])
        worksheet.append(["Кызылординская область", 13])
        worksheet.append([
            "Западно-Казахстанская область\n г. Орал\n (склад авто по взаимке)",
            395,
        ])

        table = build_table4(
            None,
            {"worksheets": [worksheet]},
            load_config().normalization,
        )
        values = {row["Точка"]: row["Количество НП"] for row in table["rows"]}

        self.assertEqual(values["Алматинская область"], 1454)
        self.assertEqual(values["Кызылординская область"], 13)
        self.assertEqual(values["г. Уральск, склад"], 395)
        self.assertNotIn("ЗКО", values)
        self.assertNotIn("Атырауская область", values)
        self.assertNotIn("Жамбылская область", values)
        self.assertNotIn("Область Жетису", values)
        self.assertNotIn("Область Улытау", values)
        self.assertNotIn("тп «Жибек жолы»", values)
        self.assertNotIn("тп «Капланбек»", values)
        self.assertEqual(
            [row["Точка"] for row in table["rows"][-2:]],
            ["Алматинская область", "Кызылординская область"],
        )
        self.assertEqual(
            sum(value for value in values.values() if isinstance(value, (int, float))),
            1454 + 13 + 395,
        )

    def test_elapsed_days_restarts_each_calendar_month(self) -> None:
        self.assertEqual(_stage2_elapsed_days(dt.date(2026, 7, 31)), 1)
        self.assertEqual(_stage2_elapsed_days(dt.date(2026, 8, 4)), 1)
        self.assertEqual(_stage2_elapsed_days(dt.date(2026, 8, 31)), 1)
        self.assertEqual(_stage2_elapsed_days(dt.date(2026, 9, 1)), 1)
        self.assertEqual(_stage2_elapsed_days(dt.date(2026, 9, 2)), 2)
        self.assertEqual(_stage2_elapsed_days(dt.date(2026, 10, 1)), 1)

    def test_duplicate_date_uses_populated_actual_column(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "СЕНТЯБРЬ"
        worksheet.append([
            "Пост",
            dt.datetime(2026, 9, 8),
            dt.datetime(2026, 9, 8),
        ])
        worksheet.append(["Алакөл (авто)", 1333, None])
        worksheet.append(["Нұржолы (авто)", 768, None])

        table = build_table4(
            None,
            {"worksheets": [worksheet]},
            load_config().normalization,
        )
        values = {row["Точка"]: row["Количество НП"] for row in table["rows"]}

        self.assertEqual(values["тп «Алаколь»"], 1333)
        self.assertEqual(values["тп «Нур Жолы»"], 768)

    def test_total_adds_karasu_and_daily_table_starts_at_h1(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        table = {
            "rows": [
                {"Точка": "тп «Тажен»", "Количество НП": 100},
                {"Точка": "тп «Бахты»", "Количество НП": 200},
            ],
            "used_np_rows": [
                {"Пункт": "тп «Тажен»", "Количество": 528},
                {"Пункт": "тп «Болашак»", "Количество": 0},
                {"Пункт": "тп «Бахты»", "Количество": "-"},
                {"Пункт": "Итого", "Количество": 528},
            ],
        }

        with patch("io_excel._stage2_elapsed_days", return_value=4):
            _write_table4(worksheet, table)

        self.assertEqual(worksheet["A4"].value, "Итого(с учетом пилот Карасу)")
        self.assertEqual(worksheet["B4"].value, f"=SUM(B2:B3)+{PILOT_KARASU_NP}")
        self.assertIsNone(worksheet["G1"].value)
        self.assertEqual(worksheet["H1"].value, "Ежедневный расход пломб")
        self.assertEqual(worksheet["H2"].value, "Количество прошедших дней текущего месяца")
        self.assertEqual(worksheet["J2"].value, 4)
        self.assertEqual(worksheet["H4"].value, "тп «Тажен»")
        self.assertEqual(worksheet["I3"].value, "Использовано НП\nв текущем месяце")
        self.assertEqual(worksheet["I4"].value, 528)
        self.assertEqual(worksheet["J4"].value, '=IF(ISNUMBER(I4),ROUND(I4/$J$2,0),"-")')
        self.assertEqual(worksheet["F3"].value, "-")
        self.assertEqual(worksheet["I5"].value, "-")
        self.assertEqual(worksheet["H7"].value, "Итого")
        self.assertEqual(worksheet["F5"].value, 528)
        self.assertEqual(worksheet["I7"].value, 528)
        self.assertEqual(worksheet["J7"].value, '=IF(ISNUMBER(I7),ROUND(I7/$J$2,0),"-")')


if __name__ == "__main__":
    unittest.main()
