from __future__ import annotations

import datetime as dt
import json
import re
from collections import Counter
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Iterable

from openpyxl import Workbook, load_workbook

from builders import build_table4


STAGE2_START = dt.datetime(2026, 9, 1, 0, 0, 0)

REQUIRED_HEADERS = {
    "Дата создания",
    "Владелец",
    "Номер перевозки",
    "Статус",
    "Дата обновления",
    "Тип перевозки",
    "Тип транспорта",
    "Компания",
    "БИН",
    "Номера документов",
    "Номера ТС",
    "Коды ТН ВЭД",
    "Номера НП",
    "Кол. путевых точек",
    "Кол. водителей",
    "Кол. операторов",
    "Начальный пункт",
    "Конечный пункт",
    "Страна начала",
    "Страна завершения",
    "Промежуточные точки",
    "Промежуточные страны",
}

# В выгрузке от 10.09.2026 несколько служебных столбцов были переименованы.
# Внутри расчёта сохраняем единый набор канонических названий, чтобы смена
# порядка/подписей колонок не меняла результаты.
HEADER_ALIASES = {
    "Водителей": "Кол. водителей",
    "Путевых точек": "Кол. путевых точек",
    "Операторов": "Кол. операторов",
    "Промежуточные пункты": "Промежуточные точки",
}

COUNTRY_CODES = {
    "российская федерация": "RU",
    "россия": "RU",
    "рф": "RU",
    "республика беларусь": "BY",
    "беларусь": "BY",
    "рб": "BY",
    "кыргызская республика": "KG",
    "кыргызстан": "KG",
    "киргизия": "KG",
    "кр": "KG",
    "республика армения": "ARM",
    "армения": "ARM",
    "ра": "ARM",
    "республика казахстан": "KZ",
    "казахстан": "KZ",
    "рк": "KZ",
}

COUNTRY_NAMES = {
    "RU": "Российская Федерация",
    "BY": "Республика Беларусь",
    "KG": "Кыргызская Республика",
    "ARM": "Республика Армения",
}

COUNTRY_COLUMNS = {
    "RU": "РФ",
    "BY": "РБ",
    "KG": "КР",
    "ARM": "РА",
}

# Порядок направлений новой Таблицы3 фиксирован пользовательским шаблоном.
# Все четыре пары проходят через Казахстан, при этом Казахстан не является
# ни страной начала, ни страной завершения перевозки.
TRACKING_DIRECTIONS = (
    ("BY", "KG", "РБ - КР"),
    ("RU", "KG", "РФ - КР"),
    ("KG", "RU", "КР - РФ"),
    ("KG", "BY", "КР - РБ"),
)
TRACKING_DIRECTION_BY_COUNTRIES = {
    (country_from, country_to): label
    for country_from, country_to, label in TRACKING_DIRECTIONS
}

OUT_TITLES = {
    "RU": "В Российскую Федерацию",
    "BY": "В Республику Беларусь",
    "KG": "В Кыргызскую Республику",
    "ARM": "В Республику Армения",
}

IN_TITLES = {
    "RU": "Из Российской Федерации",
    "BY": "Из Республики Беларусь",
    "KG": "Из Кыргызской Республики",
    "ARM": "Из Республики Армения",
}

FINISHED_STATUSES = {
    "завершенная перевозка",
    "деактивированная перевозка",
}

USED_NP_STATUSES = FINISHED_STATUSES | {"активированная перевозка"}

RAIL_REGISTERED_STATUSES = {
    "активированная перевозка",
    "завершенная перевозка",
}

COUNTED_CURRENT_STATUSES = RAIL_REGISTERED_STATUSES

ALLOWED_RAIL_OWNER_PARTS = (
    "институт космической техники и технологий",
    "кеден",
    "эсф",
)

ALLOWED_AUTO_OUT_OWNER_PARTS = (
    "институт космической техники и технологий",
    "кеден",
    "эсф",
    "интеграц",
)


T1_POINT_ALIASES: Dict[str, tuple[str, ...]] = {
    "Алаколь": ("алаколь",),
    "Казыгурт": ("казыгурт",),
    "Нур жолы": ("нур жолы", "нуржолы"),
    "Бахты": ("бахты", "бакты"),
    "Калжат": ("калжат",),
    "Тажен": ("тажен",),
    "Атамекен": ("атамекен",),
    "Б. конысбаева": ("б конысбаева", "бауыржана конысбаева"),
    "Майкапчагай": ("майкапчагай", "майкапшагай"),
    "Темир-баба": ("темир баба",),
    "Капланбек": ("капланбек",),
    "Морпорт актау": ("морпорт", "актау порт"),
    "Курык порт": ("курык порт",),
    "Курык жд": ("курык жд",),
    "Болашак": ("болашак",),
    "Оазис": ("оазис",),
    "Сарыагаш": ("сарыагаш",),
    "Алтынколь": ("алтынколь", "алтынколь жол"),
    "Темиржол": ("темиржол",),
}

USED_POINT_ALIASES: Dict[str, tuple[str, ...]] = {
    "тп «Тажен»": ("тажен",),
    "тп «Темир-Баба»": ("темир баба",),
    "тп «Болашак»": ("болашак",),
    "тп «Курык»": ("курык",),
    "тп «Морпорт»": ("морпорт", "актау порт"),
    "тп «Оазис»": ("оазис", "бейнеу"),
    "тп «Бахты»": ("бахты", "бакты"),
    "тп «Майкапчагай»": ("майкапчагай", "майкапшагай"),
    "тп «Б. Конысбаева»": ("б конысбаева", "бауыржана конысбаева"),
    "тп «Казыгурт»": ("казыгурт",),
    "тп «ст. Сарыагаш»": ("сарыагаш",),
    "тп «Атамекен»": ("атамекен",),
    "тп «Калжат»": ("калжат",),
    "тп «Алаколь»": ("алаколь",),
    "тп «Темиржол»": ("темиржол",),
    "тп «Алтынколь-жол»": ("алтынколь", "алтынколь жол"),
    "тп «Нур Жолы»": ("нур жолы", "нуржолы"),
    "г. Астана склад": ("астана",),
    "г. Караганда склад": ("караганда",),
    "г. Шымкент, склад": ("шымкент",),
    "г. Петропавловск, склад": ("петропавловск", "петропавл"),
    "г. Кокшетау, склад": ("кокшетау",),
    "г. Уральск, склад": ("уральск", "орал"),
    "г. Актобе, склад": ("актобе",),
    "г. Павлодар склад": ("павлодар",),
    "г. Оскемен, склад": ("оскемен", "усть каменогорск"),
    "г. Семей, склад": ("семей",),
    "г. Актау": ("актау",),
    "Атырауская область": ("атырау",),
    "ЗКО": ("западно казахстанская область",),
    "Костанайская область": ("костанай",),
}


def _norm(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if text == "null":
        return ""
    for old, new in (
        ("ё", "е"),
        ("ө", "о"),
        ("ұ", "у"),
        ("ү", "у"),
        ("қ", "к"),
        ("ғ", "г"),
        ("ә", "а"),
        ("і", "и"),
        ("ң", "н"),
        ("–", "-"),
        ("—", "-"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _parse_datetime(value: Any) -> dt.datetime | None:
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))
        except (OverflowError, TypeError, ValueError):
            return None
    text = str(value or "").strip()
    if not text or text.lower() == "null":
        return None
    try:
        return dt.datetime.fromisoformat(text.replace("Z", "+00:00")).replace(
            tzinfo=None
        )
    except ValueError:
        pass
    for date_format in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y",
    ):
        try:
            return dt.datetime.strptime(text, date_format)
        except ValueError:
            continue
    return None


def _country_code(value: Any) -> str | None:
    return COUNTRY_CODES.get(_norm(value))


def _transport(value: Any) -> str | None:
    text = _norm(value)
    # В исходной большой выгрузке встречается системная опечатка
    # "Железнодорожный траспорт". Проверка по устойчивой основе слова
    # намеренно распознаёт и правильное, и ошибочное написание.
    if "железнодорож" in text:
        return "rail"
    if "автомоб" in text:
        return "auto"
    return None


def build_tracking_table(
    values: Counter[tuple[str, str]] | None = None,
) -> Dict[str, Any]:
    """Собрать новую Таблицу3 по проследованию через территорию РК."""
    source = values or Counter()
    direction_labels = [item[2] for item in TRACKING_DIRECTIONS]
    rows: list[Dict[str, Any]] = []

    for transport_group, transport_label in (("auto", "Авто"), ("rail", "ЖД")):
        row: Dict[str, Any] = {"Тип транспорта": transport_label}
        for direction in direction_labels:
            row[direction] = int(source[(transport_group, direction)])
        row["Итого"] = sum(int(row[direction]) for direction in direction_labels)
        rows.append(row)

    total_row: Dict[str, Any] = {"Тип транспорта": "Итого"}
    for direction in direction_labels:
        total_row[direction] = sum(int(row[direction]) for row in rows)
    total_row["Итого"] = sum(
        int(total_row[direction]) for direction in direction_labels
    )
    rows.append(total_row)

    return {
        "table": "Таблица3",
        "stage": "stage2-direct",
        "rows": rows,
        "directions": direction_labels,
    }


@lru_cache(maxsize=4)
def _load_vehicle_plate_registry(path_text: str) -> Dict[str, Any]:
    with Path(path_text).open("r", encoding="utf-8") as stream:
        return json.load(stream)


def _vehicle_number_tokens(value: Any, separator_pattern: str) -> list[str]:
    text = str(value or "").upper().strip()
    if not text or text.lower() == "null":
        return []
    tokens: list[str] = []
    for part in re.split(separator_pattern, text):
        token = re.sub(r"[.\s]", "", part).strip()
        if token:
            tokens.append(token)
    return tokens


def detect_transport_by_vehicle_numbers(
    value: Any,
    registry_path: Path,
) -> str | None:
    """Определяет авто или ЖД по номерам ТС из большой выгрузки.

    Одна запись может содержать несколько номеров. Автомобильный номер имеет
    приоритет, поскольку такой состав учитывается в автомобильной части отчета.
    """
    registry = _load_vehicle_plate_registry(str(registry_path.resolve()))
    tokens = _vehicle_number_tokens(
        value,
        str(registry.get("token_separator_pattern") or r"[,;/|\n]+"),
    )
    if not tokens:
        return None
    for group, output in (("road_patterns", "auto"), ("rail_patterns", "rail")):
        patterns = registry.get(group) or []
        for token in tokens:
            if any(
                re.fullmatch(str(item.get("regex") or ""), token)
                for item in patterns
                if str(item.get("regex") or "")
            ):
                return output
    return None


def _transportation_type(value: Any) -> str | None:
    text = _norm(value)
    if "транзит" in text:
        return "transit"
    if "взаимной торговли" in text:
        return "trade"
    if "экспорт" in text:
        return "export"
    return None


def _np_numbers(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text or text.lower() == "null":
        return []
    return [part.strip() for part in re.split(r"[,;\n]+", text) if part.strip()]


def _header_map(values: Iterable[Any]) -> Dict[str, int]:
    """Return source-column indexes using the calculation's canonical names."""
    result: Dict[str, int] = {}
    for column_index, value in enumerate(values):
        if value is None:
            continue
        source_name = str(value).strip()
        if not source_name:
            continue
        canonical_name = HEADER_ALIASES.get(source_name, source_name)
        result.setdefault(canonical_name, column_index)
    return result


def _find_header(ws, scan_rows: int = 50) -> tuple[int, Dict[str, int]] | None:
    for row_index, values in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, scan_rows), values_only=True),
        start=1,
    ):
        header_map = _header_map(values)
        if REQUIRED_HEADERS.issubset(header_map):
            return row_index, header_map
    return None


def _is_html_table_file(path: Path) -> bool:
    """Exports with an .xls suffix may actually contain an HTML table."""
    try:
        with path.open("rb") as source:
            return source.read(512).lstrip().lower().startswith(b"<html")
    except OSError:
        return False


def _find_html_header(rows: list[list[str]], scan_rows: int = 50) -> tuple[int, Dict[str, int]] | None:
    for row_index, values in enumerate(rows[:scan_rows], start=1):
        header_map = _header_map(values)
        if REQUIRED_HEADERS.issubset(header_map):
            return row_index, header_map
    return None


def _open_big_export_source(
    path: Path,
) -> tuple[Workbook | None, int, Dict[str, int], Iterable[tuple[Any, ...] | list[str]]]:
    """Open either XLSX or the HTML-in-XLS format used by the large export."""
    if _is_html_table_file(path):
        # Share the streaming HTML reader used by RAW04.  The old parser
        # materialized every cell via HTMLParser before any calculation began.
        from daily_raw04 import _iter_html_rows

        rows = _iter_html_rows(path)
        for row_index, values in enumerate(rows, start=1):
            if row_index > 50:
                break
            header_map = _header_map(values)
            if REQUIRED_HEADERS.issubset(header_map):
                return None, row_index, header_map, rows
        raise ValueError("Не найдена строка заголовков большой выгрузки")

    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        found = _find_header(sheet)
        if found is not None:
            header_row, header_map = found
            return (
                workbook,
                header_row,
                header_map,
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
            )
    workbook.close()
    raise ValueError("Не найдена строка заголовков большой выгрузки")


def is_big_export_file(path: Path) -> bool:
    if path.suffix.lower() not in {".xlsx", ".xls"} or path.name.startswith("~$"):
        return False
    workbook = None
    try:
        workbook, _, _, _ = _open_big_export_source(path)
        return True
    except Exception:
        return False
    finally:
        if workbook is not None:
            workbook.close()


def find_big_export_file(in_dir: Path) -> Path | None:
    found = [
        path
        for path in in_dir.iterdir()
        if path.is_file() and path.suffix.lower() in {".xlsx", ".xls"}
        and is_big_export_file(path)
    ]
    if not found:
        print("Ошибка: большая выгрузка не найдена.")
        print("В папке in должен находиться один XLSX или HTML-XLS с колонкой 'Дата создания'.")
        return None
    if len(found) > 1:
        print("Ошибка: найдено несколько больших выгрузок:")
        for path in found:
            print(" -", path.name)
        return None
    print("[new method] big export FOUND:", found[0].name)
    return found[0]


def _matches_alias(text: str, alias: str) -> bool:
    if text == alias:
        return True
    return f" {alias} " in f" {text} "


def _match_point(values: Iterable[Any], aliases: Dict[str, tuple[str, ...]]) -> str | None:
    normalized_values = [_norm(value) for value in values if _norm(value)]
    for value in normalized_values:
        for output_name, variants in aliases.items():
            for variant in variants:
                if _matches_alias(value, _norm(variant)):
                    return output_name
    return None


def _flow_rows(values: Counter[tuple[str, str]], transport: str, incoming: bool) -> list[Dict[str, Any]]:
    rows: list[Dict[str, Any]] = []
    titles = IN_TITLES if incoming else OUT_TITLES
    for code in ("RU", "BY", "KG", "ARM"):
        rows.append(
            {
                "Направление": titles[code],
                "Количество": int(values[(transport, code)]),
            }
        )
    rows.append(
        {
            "Направление": "Итого",
            "Количество": sum(int(row["Количество"]) for row in rows),
        }
    )
    return rows


def _country_rows(values: Counter[str]) -> list[Dict[str, Any]]:
    rows = [
        {
            "№": index,
            "Страна": COUNTRY_NAMES[code],
            "Количество перевозок": int(values[code]),
        }
        for index, code in enumerate(("RU", "BY", "KG", "ARM"), start=1)
    ]
    rows.append(
        {
            "№": 5,
            "Страна": "Итого",
            "Количество перевозок": sum(
                int(row["Количество перевозок"]) for row in rows
            ),
        }
    )
    return rows


def _split_country_rows(
    auto_values: Counter[str],
    rail_values: Counter[str],
) -> list[Dict[str, Any]]:
    rows: list[Dict[str, Any]] = []
    for index, code in enumerate(("RU", "BY", "KG", "ARM"), start=1):
        auto_count = int(auto_values[code])
        rail_count = int(rail_values[code])
        rows.append(
            {
                "№": index,
                "Страна": COUNTRY_NAMES[code],
                "Авто": auto_count,
                "ЖД": rail_count,
                "Итого": auto_count + rail_count,
            }
        )
    rows.append(
        {
            "№": 5,
            "Страна": "Итого",
            "Авто": sum(int(row["Авто"]) for row in rows),
            "ЖД": sum(int(row["ЖД"]) for row in rows),
            "Итого": sum(int(row["Итого"]) for row in rows),
        }
    )
    return rows


def _load_point_registry(registry_path: Path) -> list[Dict[str, Any]]:
    with registry_path.open("r", encoding="utf-8") as stream:
        return json.load(stream).get("points", [])


def _point_name(point: Dict[str, Any]) -> str:
    return str(point.get("report_name") or point.get("name") or "").strip()


def _point_variants(
    point: Dict[str, Any],
    source: str | None = None,
) -> list[str]:
    variants = [
        point.get("report_name"),
        point.get("name"),
        *(point.get("aliases") or []),
    ]
    source_aliases = point.get("source_aliases") or {}
    if source is None:
        for values in source_aliases.values():
            variants.extend(values or [])
    else:
        variants.extend(source_aliases.get(source) or [])
    return [str(value).strip() for value in variants if str(value or "").strip()]


def load_t1_point_aliases(
    registry_path: Path,
    labels: Iterable[str],
    source: str | None = None,
    transport: str | None = None,
) -> Dict[str, tuple[str, ...]]:
    requested = {str(label) for label in labels}
    result: Dict[str, tuple[str, ...]] = {}
    for point in _load_point_registry(registry_path):
        name = _point_name(point)
        if not name or name not in requested:
            continue
        point_transport = {"авто": "auto", "жд": "rail"}.get(
            _norm(point.get("transport"))
        )
        if transport is not None and point_transport != transport:
            continue
        result[name] = tuple(_point_variants(point, source=source))
    return result


def load_t1_transport_groups(
    registry_path: Path,
    labels: Iterable[str],
) -> Dict[str, set[str]]:
    points = _load_point_registry(registry_path)

    result: Dict[str, set[str]] = {}
    for label in labels:
        normalized_label = _norm(label)
        groups: set[str] = set()
        for point in points:
            transport = {"авто": "auto", "жд": "rail"}.get(
                _norm(point.get("transport"))
            )
            if transport is None:
                continue
            variants = _point_variants(point)
            if any(
                _matches_alias(normalized_label, _norm(variant))
                or _matches_alias(_norm(variant), normalized_label)
                for variant in variants
                if _norm(variant)
            ):
                groups.add(transport)
        result[str(label)] = groups
    return result


def _load_location_sets(data_dir: Path) -> Dict[str, set[str]]:
    files = {
        "RU": "locations_RF.json",
        "BY": "locations_RB.json",
        "KG": "locations_KG.json",
        "KZ": "locations_KZ.json",
    }
    result: Dict[str, set[str]] = {}
    for code, filename in files.items():
        path = data_dir / filename
        if not path.exists():
            result[code] = set()
            continue
        with path.open("r", encoding="utf-8") as stream:
            locations = json.load(stream).get("locations", [])
        result[code] = {_norm(value) for value in locations if _norm(value)}
    return result


def _rail_destination_country(
    value: Any,
    location_sets: Dict[str, set[str]],
) -> str | None:
    key = _norm(value)
    if not key:
        return None
    for code in ("RU", "BY", "KG", "KZ"):
        if key in location_sets.get(code, set()):
            return code
    if (
        re.match(r"^10\d{6}\b", key)
        or any(
            part in key
            for part in ("росси", "себеж", "владивосток", "ульяновск")
        )
    ):
        return "RU"
    if any(part in key for part in ("беларус", "брест", "минск", "колядич")):
        return "BY"
    if any(part in key for part in ("кыргыз", "киргиз", "аламедин", "бишкек")):
        return "KG"
    if "армени" in key:
        return "ARM"
    if "казахстан" in key:
        return "KZ"
    return None


def build_rail_transit_by_point(
    rail_raw: Any | None,
    registry_path: Path,
    labels: Iterable[str],
    data_dir: Path,
) -> tuple[Counter[tuple[str, str]], Dict[str, int]]:
    values: Counter[tuple[str, str]] = Counter()
    stats = {
        "rows_used": 0,
        "unresolved_points": 0,
        "unresolved_destinations": 0,
    }
    ws = rail_raw.get("ws") if isinstance(rail_raw, dict) else None
    if ws is None:
        return values, stats

    header_row = None
    header_map: Dict[str, int] = {}
    required = {
        "дата",
        "владелец записи",
        "уникальный номер перевозки",
        "статус",
        "company id",
        "start waypoint",
        "end waypoint",
    }
    for row_index in range(1, ws.max_row + 1):
        normalized = {
            _norm(ws.cell(row=row_index, column=column).value): column
            for column in range(1, ws.max_column + 1)
            if _norm(ws.cell(row=row_index, column=column).value)
        }
        if required.issubset(normalized):
            header_row = row_index
            header_map = normalized
            break
    if header_row is None:
        return values, stats

    aliases = load_t1_point_aliases(
        registry_path,
        labels,
        source="rail_export",
        transport="rail",
    )
    locations = _load_location_sets(data_dir)
    seen_numbers: set[str] = set()
    for row_index in range(header_row + 1, ws.max_row + 1):
        created = _parse_datetime(
            ws.cell(row=row_index, column=header_map["дата"]).value
        )
        if created is None or created < STAGE2_START:
            continue
        owner = _norm(
            ws.cell(row=row_index, column=header_map["владелец записи"]).value
        )
        status = _norm(
            ws.cell(row=row_index, column=header_map["статус"]).value
        )
        company_id = str(
            ws.cell(row=row_index, column=header_map["company id"]).value or ""
        ).strip()
        if company_id.endswith(".0"):
            company_id = company_id[:-2]
        if (
            status not in {"activated", "finished"}
            or not any(part in owner for part in ALLOWED_RAIL_OWNER_PARTS)
            or company_id == "1"
        ):
            continue
        number = str(
            ws.cell(
                row=row_index,
                column=header_map["уникальный номер перевозки"],
            ).value
            or ""
        ).strip()
        if number and number in seen_numbers:
            continue
        if number:
            seen_numbers.add(number)

        point = _match_point(
            (ws.cell(row=row_index, column=header_map["start waypoint"]).value,),
            aliases,
        )
        if point is None:
            stats["unresolved_points"] += 1
            continue
        country = _rail_destination_country(
            ws.cell(row=row_index, column=header_map["end waypoint"]).value,
            locations,
        )
        if country not in {"RU", "BY", "KG", "ARM"}:
            stats["unresolved_destinations"] += 1
            continue
        values[(point, country)] += 1
        stats["rows_used"] += 1
    return values, stats


def build_rail_export_and_trade(
    rail_raw: Any | None,
    data_dir: Path,
) -> tuple[Counter[str], Counter[str], Dict[str, int]]:
    """Считает экспорт и взаимную торговлю только по ЖД выгрузке."""
    export_values: Counter[str] = Counter()
    trade_values: Counter[str] = Counter()
    stats = {
        "rows_used": 0,
        "unresolved_destinations": 0,
    }
    ws = rail_raw.get("ws") if isinstance(rail_raw, dict) else None
    if ws is None:
        return export_values, trade_values, stats

    header_row = None
    header_map: Dict[str, int] = {}
    required = {
        "дата",
        "владелец записи",
        "уникальный номер перевозки",
        "статус",
        "тип",
        "company id",
        "end waypoint",
    }
    for row_index in range(1, ws.max_row + 1):
        normalized = {
            _norm(ws.cell(row=row_index, column=column).value): column
            for column in range(1, ws.max_column + 1)
            if _norm(ws.cell(row=row_index, column=column).value)
        }
        if required.issubset(normalized):
            header_row = row_index
            header_map = normalized
            break
    if header_row is None:
        return export_values, trade_values, stats

    locations = _load_location_sets(data_dir)
    records: Dict[str, tuple[dt.datetime, int, str, str, Any]] = {}
    for row_index in range(header_row + 1, ws.max_row + 1):
        created = _parse_datetime(
            ws.cell(row=row_index, column=header_map["дата"]).value
        )
        if created is None or created < STAGE2_START:
            continue
        owner = _norm(
            ws.cell(row=row_index, column=header_map["владелец записи"]).value
        )
        company_id = str(
            ws.cell(row=row_index, column=header_map["company id"]).value or ""
        ).strip()
        if company_id.endswith(".0"):
            company_id = company_id[:-2]
        if (
            not any(part in owner for part in ALLOWED_RAIL_OWNER_PARTS)
            or company_id == "1"
        ):
            continue
        number = str(
            ws.cell(
                row=row_index,
                column=header_map["уникальный номер перевозки"],
            ).value
            or ""
        ).strip()
        record_key = number or f"__row_{row_index}"
        status = _norm(ws.cell(row=row_index, column=header_map["статус"]).value)
        transportation_type = _transportation_type(
            ws.cell(row=row_index, column=header_map["тип"]).value
        )
        end_point = ws.cell(row=row_index, column=header_map["end waypoint"]).value
        current = records.get(record_key)
        candidate = (created, row_index, status, transportation_type or "", end_point)
        if current is None or candidate[:2] > current[:2]:
            records[record_key] = candidate

    for _, (_, _, status, transportation_type, end_point) in records.items():
        if status not in {"activated", "finished"}:
            continue
        if transportation_type not in {"export", "trade"}:
            continue
        country = _rail_destination_country(end_point, locations)
        if country not in {"RU", "BY", "KG", "ARM"}:
            stats["unresolved_destinations"] += 1
            continue
        if transportation_type == "export":
            export_values[country] += 1
        else:
            trade_values[country] += 1
        stats["rows_used"] += 1
    return export_values, trade_values, stats


def extract_stage2_rail_ids(rail_raw: Any | None) -> set[str]:
    if rail_raw is None:
        return set()
    ws = rail_raw.get("ws") if isinstance(rail_raw, dict) else None
    if ws is None:
        return set()

    header_row = None
    header_map: Dict[str, int] = {}
    for row_index in range(1, ws.max_row + 1):
        values = [
            ws.cell(row=row_index, column=column).value
            for column in range(1, ws.max_column + 1)
        ]
        normalized = {_norm(value): column for column, value in enumerate(values, 1)}
        if {
            "дата",
            "владелец записи",
            "уникальный номер перевозки",
            "статус",
            "company id",
        }.issubset(normalized):
            header_row = row_index
            header_map = normalized
            break
    if header_row is None:
        return set()

    allowed_owners = (
        "институт космической техники и технологий",
        "кеден",
        "эсф",
    )
    result: set[str] = set()
    for row_index in range(header_row + 1, ws.max_row + 1):
        created = _parse_datetime(
            ws.cell(row=row_index, column=header_map["дата"]).value
        )
        if created is None or created < STAGE2_START:
            continue
        owner = _norm(
            ws.cell(row=row_index, column=header_map["владелец записи"]).value
        )
        status = _norm(ws.cell(row=row_index, column=header_map["статус"]).value)
        company_id_value = ws.cell(
            row=row_index,
            column=header_map["company id"],
        ).value
        company_id = str(company_id_value or "").strip()
        if company_id.endswith(".0"):
            company_id = company_id[:-2]
        if (
            status not in {"activated", "finished"}
            or not any(part in owner for part in allowed_owners)
            or company_id == "1"
        ):
            continue
        transportation_number = str(
            ws.cell(
                row=row_index,
                column=header_map["уникальный номер перевозки"],
            ).value
            or ""
        ).strip()
        if transportation_number:
            result.add(transportation_number)
    return result


def _used_np_rows(labels: Iterable[str], values: Counter[str]) -> list[Dict[str, Any]]:
    rows: list[Dict[str, Any]] = []
    total = 0
    for label in labels:
        if label == "Итого":
            continue
        value = int(values[label])
        total += value
        rows.append(
            {
                "Пункт": label,
                "Количество": value if value else "-",
            }
        )
    rows.append({"Пункт": "Итого", "Количество": total})
    return rows


def _seal_post_rows(
    values: Counter[tuple[str, str, str]],
    labels: Iterable[str],
    transport: str,
) -> list[Dict[str, Any]]:
    """Build the fixed post order for the seal-application table."""
    rows: list[Dict[str, Any]] = []
    for point in labels:
        row: Dict[str, Any] = {"ТП/ЖД": point}
        for code, column in COUNTRY_COLUMNS.items():
            row[column] = int(values[(transport, point, code)])
        row["Итого"] = sum(int(row[column]) for column in COUNTRY_COLUMNS.values())
        rows.append(row)
    total: Dict[str, Any] = {"ТП/ЖД": "Итого"}
    for column in COUNTRY_COLUMNS.values():
        total[column] = sum(int(row[column]) for row in rows)
    total["Итого"] = sum(int(total[column]) for column in COUNTRY_COLUMNS.values())
    rows.append(total)
    return rows


def build_tables_from_big_export(
    path: Path,
    table4_source: Any | None,
    cfg: Dict[str, Any],
    report_date: dt.date | None = None,
    rail_transport_ids: set[str] | None = None,
    split_by_transport: bool = False,
    rail_raw: Any | None = None,
    point_registry_path: Path | None = None,
    data_dir: Path | None = None,
    vehicle_registry_path: Path | None = None,
    include_big_export_table4: bool = True,
    period_start: dt.datetime = STAGE2_START,
) -> list[Dict[str, Any]]:
    report_date = report_date or dt.date.today()
    target_codes = {"RU", "BY", "KG", "ARM"}

    t1_values: Counter[tuple[str, str]] = Counter()
    export_values: Counter[str] = Counter()
    trade_values: Counter[str] = Counter()
    split_t1_values: Counter[tuple[str, str, str]] = Counter()
    split_export_values = {"auto": Counter(), "rail": Counter()}
    split_trade_values = {"auto": Counter(), "rail": Counter()}
    tracking_values: Counter[tuple[str, str]] = Counter()
    completed_out: Counter[tuple[str, str]] = Counter()
    completed_in: Counter[tuple[str, str]] = Counter()
    active_in: Counter[str] = Counter()
    used_np_values: Counter[str] = Counter()
    monthly_np_values: Counter[str] = Counter()
    seal_post_values: Counter[tuple[str, str, str]] = Counter()
    rail_registered_total = 0
    unresolved_t1 = 0
    unresolved_np = 0
    unresolved_seal_points = 0
    invalid_dates = 0
    rows_read = 0
    duplicate_rows = 0
    # Разделение вывода авто/ЖД может работать исключительно по большой
    # выгрузке. Наличие `rail_transport_ids` сохраняет прежний режим, где
    # ЖД-перевозки дополнительно подтверждались внешней ЖД-выгрузкой.
    split_mode = split_by_transport or rail_transport_ids is not None
    rail_ids = rail_transport_ids or set()
    t1_transport_groups = cfg.get("t1_transport_groups", {}) or {}
    t1_point_aliases = T1_POINT_ALIASES
    rail_t1_values: Counter[tuple[str, str]] | None = None
    rail_t1_stats: Dict[str, int] = {}
    rail_export_values: Counter[str] | None = None
    rail_trade_values: Counter[str] | None = None
    rail_flow_stats: Dict[str, int] = {}
    vehicle_stats: Counter[str] = Counter()
    if vehicle_registry_path is None:
        vehicle_registry_path = Path(__file__).resolve().parent / "data" / "vehicle_plate_formats.json"
    if point_registry_path is not None:
        t1_point_aliases = load_t1_point_aliases(
            point_registry_path,
            cfg.get("tp_zhd_list_t1", []),
            source="big_export",
        )
    if (
        split_mode
        and rail_raw is not None
        and point_registry_path is not None
        and data_dir is not None
    ):
        rail_t1_values, rail_t1_stats = build_rail_transit_by_point(
            rail_raw,
            point_registry_path,
            cfg.get("tp_zhd_list_t1", []),
            data_dir,
        )
        rail_export_values, rail_trade_values, rail_flow_stats = (
            build_rail_export_and_trade(rail_raw, data_dir)
        )

    workbook: Workbook | None = None
    try:
        workbook, header_row, header_map, source_rows = _open_big_export_source(path)

        indexes = {name: header_map[name] for name in REQUIRED_HEADERS}
        used_labels: list[str] | None = None

        if include_big_export_table4:
            blank_workbook = Workbook()
            try:
                table5_np = build_table4(
                    {"ws": blank_workbook.active},
                    table4_source,
                    cfg,
                )
            finally:
                blank_workbook.close()
        else:
            # Режим 3 заменяет этот объект Т5 готовой таблицей из report и
            # файла готовности АПП/ЖД. Не читаем их здесь и не смешиваем с
            # номерами НП большой выгрузки.
            table5_np = {"table": "Таблица5", "used_np_rows": []}

        used_labels = [
            str(row.get("Пункт", "")).strip()
            for row in table5_np.get("used_np_rows", [])
            if str(row.get("Пункт", "")).strip()
        ]

        records: Dict[str, Dict[str, Any]] = {}
        record_np_numbers: Dict[str, set[str]] = {}

        for source_row, values in enumerate(
            source_rows,
            start=header_row + 1,
        ):
            if not any(value not in (None, "") for value in values):
                continue
            row = {
                name: values[index] if index < len(values) else None
                for name, index in indexes.items()
            }
            rows_read += 1
            if rows_read % 25_000 == 0:
                print(f"[new method] прочитано строк: {rows_read}; уникальных перевозок: {len(records)}")

            created = _parse_datetime(row["Дата создания"])
            updated = _parse_datetime(row["Дата обновления"])
            if created is None:
                invalid_dates += 1

            transportation_number = str(row["Номер перевозки"] or "").strip()
            record_key = transportation_number or f"__row_{source_row}"
            record_np_numbers.setdefault(record_key, set()).update(
                _np_numbers(row["Номера НП"])
            )
            row["_created"] = created
            row["_updated"] = updated
            row["_source_row"] = source_row
            current = records.get(record_key)
            row_order = (updated or created or dt.datetime.min, source_row)
            if current is None:
                records[record_key] = row
            else:
                duplicate_rows += 1
                current_order = (
                    current["_updated"] or current["_created"] or dt.datetime.min,
                    current["_source_row"],
                )
                if row_order > current_order:
                    records[record_key] = row

        for record_key, row in records.items():
            created = row["_created"]
            updated = row["_updated"]

            status = _norm(row["Статус"])
            transportation_type = _transportation_type(row["Тип перевозки"])
            declared_transport = _transport(row["Тип транспорта"])
            detected_transport = detect_transport_by_vehicle_numbers(
                row["Номера ТС"],
                vehicle_registry_path,
            )
            if detected_transport is not None:
                transport_group = detected_transport
                vehicle_stats[f"detected_{detected_transport}"] += 1
            elif record_key in rail_ids:
                # Подтвержденная в ЖД выгрузке перевозка не может попасть
                # в авто только из-за пустого или нестандартного номера.
                transport_group = "rail"
                vehicle_stats["confirmed_by_rail_export"] += 1
            else:
                transport_group = declared_transport or "auto"
                vehicle_stats["fallback_declared_type"] += 1
            country_from = _country_code(row["Страна начала"])
            country_to = _country_code(row["Страна завершения"])

            created_in_period = created is not None and created >= period_start
            tracking_direction = TRACKING_DIRECTION_BY_COUNTRIES.get(
                (country_from, country_to)
            )
            if (
                created_in_period
                and status in COUNTED_CURRENT_STATUSES
                and transportation_type == "transit"
                and tracking_direction is not None
            ):
                tracking_values[(transport_group, tracking_direction)] += 1

            if (
                created_in_period
                and status in COUNTED_CURRENT_STATUSES
                and transportation_type == "transit"
                and country_from == "KZ"
                and country_to in target_codes
            ):
                point = _match_point((row["Начальный пункт"],), t1_point_aliases)
                if point is None:
                    unresolved_t1 += 1
                else:
                    t1_values[(point, country_to)] += 1
                    point_groups = t1_transport_groups.get(
                        point, {"auto", "rail"}
                    )
                    if point_groups == {"auto"}:
                        t1_group = "auto"
                    elif point_groups == {"rail"}:
                        t1_group = "rail"
                    else:
                        t1_group = transport_group
                    split_t1_values[(t1_group, point, country_to)] += 1

            if (
                created_in_period
                and status in COUNTED_CURRENT_STATUSES
                and country_from == "KZ"
                and country_to in target_codes
            ):
                if transportation_type == "export":
                    export_values[country_to] += 1
                    split_export_values[transport_group][country_to] += 1
                elif transportation_type == "trade":
                    trade_values[country_to] += 1
                    split_trade_values[transport_group][country_to] += 1

            owner = _norm(row["Владелец"])
            # Таблица4 — завершенные перевозки. В отличие от старого
            # операторского фильтра, здесь тип транспорта берется из самой
            # перевозки (с проверкой номера ТС), а не из владельца записи.
            # Так в авто и ЖД попадают все завершенные перевозки из РК и в РК.
            if split_mode and (
                created_in_period
                and status == "завершенная перевозка"
                and country_from == "KZ"
                and country_to in target_codes
            ):
                completed_out[(transport_group, country_to)] += 1
            elif not split_mode and (
                created_in_period
                and status == "завершенная перевозка"
                and country_from == "KZ"
                and country_to in target_codes
            ):
                completed_out[("auto", country_to)] += 1

            if split_mode and (
                created_in_period
                and status == "завершенная перевозка"
                and country_to == "KZ"
                and country_from in target_codes
            ):
                completed_in[(transport_group, country_from)] += 1
            elif not split_mode and (
                created_in_period
                and status == "завершенная перевозка"
                and declared_transport in {"auto", "rail"}
            ):
                if (
                    declared_transport == "rail"
                    and country_from == "KZ"
                    and country_to in target_codes
                ):
                    completed_out[("rail", country_to)] += 1
                elif country_to == "KZ" and country_from in target_codes:
                    completed_in[(declared_transport, country_from)] += 1

            if (
                created_in_period
                and status == "активированная перевозка"
                and country_to == "KZ"
                and country_from in target_codes
            ):
                active_in[country_from] += 1

            if (
                created_in_period
                and transport_group == "rail"
                and status in RAIL_REGISTERED_STATUSES
                and country_from == "KZ"
                and any(part in owner for part in ALLOWED_RAIL_OWNER_PARTS)
            ):
                rail_registered_total += 1

            np_numbers = record_np_numbers[record_key]
            if (
                np_numbers
                and status in USED_NP_STATUSES
                and created_in_period
                and country_from == "KZ"
            ):
                point = _match_point((row["Начальный пункт"],), USED_POINT_ALIASES)
                if point is None:
                    point = "Навешивания на территории РК"
                    unresolved_np += len(np_numbers)
                used_np_values[point] += len(np_numbers)

                if updated is not None and (
                    updated.year,
                    updated.month,
                ) == (report_date.year, report_date.month):
                    monthly_np_values[point] += len(np_numbers)

            # Таблица7: наложенные НП по таможенным постам. Популяция
            # начинается с переданного периода, только из РК, только активные
            # и завершенные перевозки. Один номер НП считается один раз.
            if (
                np_numbers
                and status in RAIL_REGISTERED_STATUSES
                and created_in_period
                and country_from == "KZ"
            ):
                seal_point = _match_point((row["Начальный пункт"],), t1_point_aliases)
                if seal_point is None or seal_point not in t1_transport_groups:
                    unresolved_seal_points += 1
                else:
                    point_groups = t1_transport_groups.get(seal_point, {transport_group})
                    if point_groups == {"auto"}:
                        seal_transport = "auto"
                    elif point_groups == {"rail"}:
                        seal_transport = "rail"
                    else:
                        seal_transport = transport_group
                    seal_post_values[(seal_transport, seal_point, country_to)] += len(np_numbers)
    finally:
        if workbook is not None:
            workbook.close()

    t1_rows: list[Dict[str, Any]] = []
    for point in cfg.get("tp_zhd_list_t1", []):
        row: Dict[str, Any] = {"ТП/ЖД": point}
        for code, column in COUNTRY_COLUMNS.items():
            row[column] = int(t1_values[(point, code)])
        row["Итого"] = sum(int(row[column]) for column in COUNTRY_COLUMNS.values())
        t1_rows.append(row)
    summary: Dict[str, Any] = {"ТП/ЖД": "Итого"}
    for column in COUNTRY_COLUMNS.values():
        summary[column] = sum(int(row[column]) for row in t1_rows)
    summary["Итого"] = sum(int(summary[column]) for column in COUNTRY_COLUMNS.values())
    t1_rows.append(summary)

    table_stage = (
        "stage2-direct" if period_start == STAGE2_START else "full-period"
    )
    table1 = {
        "table": "Таблица1",
        "stage": table_stage,
        "rows": t1_rows,
        "rail_registered_total": rail_registered_total,
    }
    table2 = {
        "table": "Таблица2",
        "stage": table_stage,
        "export_rows": _country_rows(export_values),
        "trade_rows": _country_rows(trade_values),
    }
    if split_mode:
        def split_t1_rows(group: str) -> list[Dict[str, Any]]:
            output_rows: list[Dict[str, Any]] = []
            source_values = (
                rail_t1_values
                if group == "rail" and rail_t1_values is not None
                else split_t1_values
            )
            for point in cfg.get("tp_zhd_list_t1", []):
                if group not in t1_transport_groups.get(point, set()):
                    continue
                output_row: Dict[str, Any] = {"ТП/ЖД": point}
                for code, column in COUNTRY_COLUMNS.items():
                    if source_values is rail_t1_values:
                        output_row[column] = int(source_values[(point, code)])
                    else:
                        output_row[column] = int(source_values[(group, point, code)])
                output_row["Итого"] = sum(
                    int(output_row[column])
                    for column in COUNTRY_COLUMNS.values()
                )
                output_rows.append(output_row)
            total_row: Dict[str, Any] = {"ТП/ЖД": "Итого"}
            for column in COUNTRY_COLUMNS.values():
                total_row[column] = sum(
                    int(output_row[column]) for output_row in output_rows
                )
            total_row["Итого"] = sum(
                int(total_row[column]) for column in COUNTRY_COLUMNS.values()
            )
            output_rows.append(total_row)
            return output_rows

        table1["transport_split"] = True
        table1["auto_rows"] = split_t1_rows("auto")
        table1["rail_rows"] = split_t1_rows("rail")
        # В комбинированном старом варианте число ЖД подтверждалось отдельной
        # ЖД-выгрузкой. В новом варианте с split_by_transport оно уже посчитано
        # выше из большой выгрузки: активные и завершённые, начало в РК.
        if rail_transport_ids is not None:
            table1["rail_registered_total"] = len(rail_ids)
        table1["rail_transit_source"] = (
            "rail_export" if rail_t1_values is not None else "big_export"
        )
        table1["rail_transit_stats"] = rail_t1_stats
        table2["transport_split"] = True
        table2["trade_split_rows"] = _split_country_rows(
            split_trade_values["auto"],
            rail_trade_values
            if rail_trade_values is not None
            else split_trade_values["rail"],
        )
        table2["export_split_rows"] = _split_country_rows(
            split_export_values["auto"],
            rail_export_values
            if rail_export_values is not None
            else split_export_values["rail"],
        )
        table2["rail_source"] = (
            "rail_export" if rail_export_values is not None else "big_export"
        )
        table2["rail_stats"] = rail_flow_stats
    auto_out_rows = _flow_rows(completed_out, "auto", incoming=False)
    auto_in_rows = _flow_rows(completed_in, "auto", incoming=True)
    rail_out_rows = _flow_rows(completed_out, "rail", incoming=False)
    rail_in_rows = _flow_rows(completed_in, "rail", incoming=True)
    table3 = build_tracking_table(tracking_values)
    table3["stage"] = table_stage

    table4 = {
        "table": "Таблица4",
        "stage": table_stage,
        "auto_out_rows": auto_out_rows,
        "auto_in_rows": auto_in_rows,
        "rail_out_rows": rail_out_rows,
        "rail_in_rows": rail_in_rows,
        "auto_out_total": auto_out_rows[-1]["Количество"],
        "auto_in_total": auto_in_rows[-1]["Количество"],
        "rail_out_total": rail_out_rows[-1]["Количество"],
        "rail_in_total": rail_in_rows[-1]["Количество"],
    }

    table5_np["table"] = "Таблица5"
    table5_np["stage"] = table_stage
    table5_np["used_np_rows"] = _used_np_rows(used_labels or [], used_np_values)
    table5_np["used_np_total"] = table5_np["used_np_rows"][-1]["Количество"]
    table5_np["monthly_used_np_rows"] = _used_np_rows(
        used_labels or [], monthly_np_values
    )
    table5_np["monthly_report_date"] = report_date.isoformat()
    table5_np["monthly_elapsed_days"] = report_date.day

    table5_rows = [
        {
            "Страна начала перевозки": COUNTRY_NAMES[code],
            "Количество перевозок": int(active_in[code]),
        }
        for code in ("RU", "BY", "KG", "ARM")
    ]
    table5_rows.append(
        {
            "Страна начала перевозки": "Итого",
            "Количество перевозок": sum(
                int(row["Количество перевозок"]) for row in table5_rows
            ),
        }
    )
    table6 = {
        "table": "Таблица6",
        "stage": table_stage if table_stage == "full-period" else "current-snapshot",
        "rows": table5_rows,
    }
    auto_post_labels = [
        point for point in cfg.get("tp_zhd_list_t1", [])
        if "auto" in t1_transport_groups.get(point, set())
    ]
    rail_post_labels = [
        point for point in cfg.get("tp_zhd_list_t1", [])
        if "rail" in t1_transport_groups.get(point, set())
    ]
    auto_seal_rows = _seal_post_rows(seal_post_values, auto_post_labels, "auto")
    rail_seal_rows = _seal_post_rows(seal_post_values, rail_post_labels, "rail")
    table7 = {
        "table": "Таблица7",
        "stage": table_stage,
        "auto_rows": auto_seal_rows,
        "rail_rows": rail_seal_rows,
        "auto_total": auto_seal_rows[-1]["Итого"],
        "rail_total": rail_seal_rows[-1]["Итого"],
        "total": sum(seal_post_values.values()),
        "unresolved_points": unresolved_seal_points,
        "period_start": period_start.isoformat(),
    }

    print(
        "[new method] rows_read=", rows_read,
        " unique_transportations=", len(records),
        " duplicate_rows=", duplicate_rows,
        " invalid_creation_dates=", invalid_dates,
        " unresolved_t1=", unresolved_t1,
        " unresolved_np=", unresolved_np,
        " unresolved_seal_points=", unresolved_seal_points,
        " vehicle_auto=", vehicle_stats["detected_auto"],
        " vehicle_rail=", vehicle_stats["detected_rail"],
        " vehicle_fallback=", vehicle_stats["fallback_declared_type"],
        sep="",
    )
    print(
        "[new method] calculation period: ",
        f"from={period_start:%Y-%m-%d %H:%M:%S} ",
        "to=latest source record",
        sep="",
    )
    print(
        "[new method] totals: ",
        f"table1={summary['Итого']}, ",
        f"export={table2['export_rows'][-1]['Количество перевозок']}, ",
        f"trade={table2['trade_rows'][-1]['Количество перевозок']}, ",
        f"tracking={table3['rows'][-1]['Итого']}, ",
        f"used_np={table5_np['used_np_total']}, ",
        f"seal_post_total={table7['total']}",
        sep="",
    )
    return [table1, table2, table3, table4, table5_np, table6, table7]
