"""Daily NP report calculated from one RAW-04 transportation export.

This module is deliberately separate from the legacy/stage-2 builders.  It
implements the calculation contract used by the daily certificate and must not
change the results of the existing menu modes.
"""
from __future__ import annotations

import datetime as dt
import re
from collections import Counter, defaultdict
from html import unescape
from pathlib import Path
from typing import Any, Iterable, Iterator

from openpyxl import load_workbook

from big_export import (
    COUNTRY_COLUMNS,
    COUNTRY_NAMES,
    TRACKING_DIRECTION_BY_COUNTRIES,
    _country_code,
    _match_point,
    _norm,
    _np_numbers,
    _parse_datetime,
    _transport,
    _transportation_type,
    load_t1_point_aliases,
    load_t1_transport_groups,
)


PERIOD_START = dt.datetime(2026, 2, 11)
COUNTRIES = ("RU", "BY", "KG", "ARM")
FINISHED_STATUS = "завершенная перевозка"
APPLICATION_FALLBACK_STATUSES = {
    "активированная перевозка",
    "завершенная перевозка",
    "деактивированная перевозка",
}
COUNTED_APPLICATION_STATUSES = {
    "активированная перевозка",
    "завершенная перевозка",
}
PILOT_DIRECTION_COUNTS = {
    "Транзит": 35,
    "Экспорт": 19,
    "Взаимная торговля": 537,
}
PILOT_DIRECTION_TOTAL = sum(PILOT_DIRECTION_COUNTS.values())

# Canonical fields and acceptable normalized column names.  The raw export has
# changed labels a few times, so it is intentionally read by header meaning.
HEADER_VARIANTS = {
    "shipment": ("номер перевозки", "shipment id"),
    "updated": ("дата обновления", "дата изменения", "дата обновления перевозки"),
    "status": ("статус", "status"),
    "status_time": ("время статуса", "status date", "дата статуса"),
    "created": ("дата создания", "дата начала"),
    "activated": ("дата активации", "activation date"),
    "country_from": ("страна начала", "страна отправления"),
    "country_to": ("страна завершения", "страна назначения"),
    "intermediate_countries": ("промежуточные страны",),
    "start_point": ("начальный пункт", "start waypoint"),
    "transport_type": ("тип перевозки", "вид перевозки"),
    "transport": ("тип транспорта", "транспорт"),
    "seals": ("номера нп", "номера навигационных пломб", "seal numbers"),
}
REQUIRED = {
    "shipment", "status", "created", "country_from", "country_to",
    "start_point", "transport_type", "transport", "seals",
}


def _header_key(value: Any) -> str:
    return _norm(value)


def _header_map(values: Iterable[Any]) -> dict[str, int]:
    normalized = {_header_key(value): index for index, value in enumerate(values)}
    result: dict[str, int] = {}
    for name, variants in HEADER_VARIANTS.items():
        for variant in variants:
            index = normalized.get(_header_key(variant))
            if index is not None:
                result[name] = index
                break
    return result


def _is_html(path: Path) -> bool:
    try:
        with path.open("rb") as source:
            return source.read(512).lstrip().lower().startswith(b"<html")
    except OSError:
        return False


def _iter_html_rows(path: Path) -> Iterator[list[str]]:
    """Fast row iterator for the regular HTML table used by RAW04 exports.

    ``HTMLParser`` invokes Python callbacks for every one of roughly three
    million cells in a typical export and was the main source of the 10-minute
    wait.  Regex scanning here stays in C for table boundaries and never keeps
    a second list of all source rows in memory.
    """
    row_pattern = re.compile(rb"<tr[^>]*>(.*?)</tr\s*>", re.IGNORECASE | re.DOTALL)
    cell_pattern = re.compile(rb"<t[dh][^>]*>(.*?)</t[dh]\s*>", re.IGNORECASE | re.DOTALL)
    tag_pattern = re.compile(rb"<[^>]+>")
    raw = path.read_bytes()
    for raw_row in row_pattern.finditer(raw):
        row: list[str] = []
        for raw_cell in cell_pattern.findall(raw_row.group(1)):
            text = tag_pattern.sub(b" ", raw_cell).decode("utf-8", errors="replace")
            row.append(" ".join(unescape(text).split()))
        if row:
            yield row


def _open_source(path: Path):
    """Return a header map and source rows for XLSX and HTML-in-XLS exports."""
    if _is_html(path):
        rows = _iter_html_rows(path)
        for index, values in enumerate(rows, start=1):
            if index > 50:
                break
            mapping = _header_map(values)
            if REQUIRED.issubset(mapping):
                return None, mapping, rows
        raise ValueError("Не найдена строка обязательных заголовков RAW-04")

    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        for row_index, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 50), values_only=True),
            start=1,
        ):
            mapping = _header_map(values)
            if REQUIRED.issubset(mapping):
                return workbook, mapping, sheet.iter_rows(min_row=row_index + 1, values_only=True)
    workbook.close()
    raise ValueError("Не найдена строка обязательных заголовков RAW-04")


def is_raw04_export(path: Path) -> bool:
    if path.suffix.lower() not in {".xlsx", ".xls"} or path.name.startswith("~$"):
        return False
    # Do not parse an entire 100+ MB HTML-XLS merely to recognize it.  The
    # header is in the export preamble; the full parsing happens only once in
    # the actual calculation.
    if _is_html(path):
        try:
            with path.open("rb") as source:
                preview = source.read(2_097_152).decode("utf-8", errors="replace")
            preview_key = _norm(preview)
            return all(_header_key(value) in preview_key for value in (
                "Номер перевозки", "Страна начала", "Начальный пункт", "Номера НП",
            ))
        except OSError:
            return False
    workbook = None
    try:
        workbook, _, _ = _open_source(path)
        return True
    except Exception:
        return False
    finally:
        if workbook is not None:
            workbook.close()


def find_raw04_export(in_dir: Path) -> Path | None:
    found = [path for path in in_dir.iterdir() if path.is_file() and is_raw04_export(path)]
    if len(found) == 1:
        print("[raw04] export FOUND:", found[0].name)
        return found[0]
    if not found:
        print("Ошибка: RAW-04 выгрузка не найдена.")
    else:
        print("Ошибка: найдено несколько RAW-04 выгрузок:")
        for path in found:
            print(" -", path.name)
    return None


def _value(values: tuple[Any, ...] | list[Any], mapping: dict[str, int], name: str) -> Any:
    index = mapping.get(name)
    return values[index] if index is not None and index < len(values) else None


def _application_date(record: dict[str, Any], diagnostics: Counter[str]) -> dt.datetime | None:
    activated = record["activated"]
    if activated is not None:
        return activated
    if record["status"] in APPLICATION_FALLBACK_STATUSES and record["created"] is not None:
        diagnostics["activation_date_replaced_by_created"] += 1
        return record["created"]
    return None


def _flow_rows(values: Counter[tuple[str, str]], transport: str, incoming: bool) -> list[dict[str, Any]]:
    title = (
        {"RU": "Из Российской Федерации", "BY": "Из Республики Беларусь", "KG": "Из Кыргызской Республики", "ARM": "Из Республики Армения"}
        if incoming else
        {"RU": "В Российскую Федерацию", "BY": "В Республику Беларусь", "KG": "В Кыргызскую Республику", "ARM": "В Республику Армения"}
    )
    flow = "in" if incoming else "out"
    rows = [
        {"Направление": title[code], "Количество": int(values[(transport, flow, code)])}
        for code in COUNTRIES
    ]
    rows.append({"Направление": "Итого", "Количество": sum(item["Количество"] for item in rows)})
    return rows


def _country_rows(values: Counter[str]) -> list[dict[str, Any]]:
    rows = [{"Страна начала перевозки": COUNTRY_NAMES[code], "Количество перевозок": int(values[code])} for code in COUNTRIES]
    rows.append({"Страна начала перевозки": "Итого", "Количество перевозок": sum(item["Количество перевозок"] for item in rows)})
    return rows


def _split_country_rows(values: Counter[tuple[str, str]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for number, code in enumerate(COUNTRIES, start=1):
        auto, rail = int(values[("auto", code)]), int(values[("rail", code)])
        rows.append({"№": number, "Страна": COUNTRY_NAMES[code], "Авто": auto, "ЖД": rail, "Итого": auto + rail})
    rows.append({"№": 5, "Страна": "Итого", "Авто": sum(item["Авто"] for item in rows), "ЖД": sum(item["ЖД"] for item in rows), "Итого": sum(item["Итого"] for item in rows)})
    return rows


def _post_rows(values: Counter[tuple[str, str, str]], posts: list[str], group: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for post in posts:
        row = {"ТП/ЖД": post}
        for code, column in COUNTRY_COLUMNS.items():
            row[column] = int(values[(group, post, code)])
        row["Итого"] = sum(row[column] for column in COUNTRY_COLUMNS.values())
        rows.append(row)
    total = {"ТП/ЖД": "Итого"}
    for column in COUNTRY_COLUMNS.values():
        total[column] = sum(row[column] for row in rows)
    total["Итого"] = sum(total[column] for column in COUNTRY_COLUMNS.values())
    rows.append(total)
    return rows


def build_daily_raw04_tables(path: Path, registry_path: Path) -> list[dict[str, Any]]:
    """Build the seven certificate tables plus a technical diagnostics sheet."""
    print("[raw04] Шаг 1 из 3: читается и очищается большая RAW04-выгрузка")
    workbook, mapping, source_rows = _open_source(path)
    diagnostics: Counter[str] = Counter()
    examples: dict[str, list[str]] = defaultdict(list)
    records: dict[str, dict[str, Any]] = {}
    missing_shipment = 0
    source_count = 0
    try:
        for source_row, values in enumerate(source_rows, start=1):
            if not any(value not in (None, "") for value in values):
                continue
            source_count += 1
            shipment = str(_value(values, mapping, "shipment") or "").strip()
            key = shipment or f"__row_{source_row}"
            if not shipment:
                missing_shipment += 1
            updated = _parse_datetime(_value(values, mapping, "updated"))
            status_time = _parse_datetime(_value(values, mapping, "status_time"))
            record = {
                "shipment": shipment or f"Строка {source_row}",
                "source_row": source_row,
                "updated": updated,
                "status_time": status_time,
                "created": _parse_datetime(_value(values, mapping, "created")),
                "activated": _parse_datetime(_value(values, mapping, "activated")),
                "status": _norm(_value(values, mapping, "status")),
                "country_from": _country_code(_value(values, mapping, "country_from")),
                "country_to": _country_code(_value(values, mapping, "country_to")),
                "intermediate": _norm(_value(values, mapping, "intermediate_countries")),
                "point_raw": _value(values, mapping, "start_point"),
                "transport_type": _transportation_type(_value(values, mapping, "transport_type")),
                "transport": _transport(_value(values, mapping, "transport")),
                "np_numbers": set(_np_numbers(_value(values, mapping, "seals"))),
            }
            current = records.get(key)
            order = (updated or status_time or dt.datetime.min, source_row)
            if current is None or order >= (current["updated"] or current["status_time"] or dt.datetime.min, current["source_row"]):
                records[key] = record
            if source_row % 25_000 == 0:
                print(
                    "[raw04] прочитано строк: ", source_row,
                    "; актуальных перевозок: ", len(records),
                    sep="",
                )
    finally:
        if workbook is not None:
            workbook.close()

    print(
        "[raw04] Шаг 1 завершен: строк=", source_count,
        ", актуальных перевозок=", len(records),
        sep="",
    )

    point_labels = [
        "Алаколь", "Казыгурт", "Нур жолы", "Калжат", "Бахты", "Тажен",
        "Атамекен", "Б. конысбаева", "Майкапчагай", "Темир-баба", "Капланбек",
        "Морпорт актау", "Курык порт", "Болашак", "Оазис", "Сарыагаш", "Алтынколь", "Темиржол",
    ]
    aliases = load_t1_point_aliases(registry_path, point_labels)
    groups = load_t1_transport_groups(registry_path, point_labels)
    auto_posts = [name for name in point_labels if groups.get(name) == {"auto"}]
    rail_posts = [name for name in point_labels if groups.get(name) == {"rail"}]

    transit: Counter[tuple[str, str, str]] = Counter()
    export: Counter[tuple[str, str]] = Counter()
    mutual: Counter[tuple[str, str]] = Counter()
    tracking: Counter[tuple[str, str]] = Counter()
    completed: Counter[tuple[str, str, str]] = Counter()
    completed_in_all: Counter[str] = Counter()
    applications: Counter[tuple[str, str]] = Counter()
    post_applications: Counter[tuple[str, str, str]] = Counter()
    unknown_post_pairs = 0
    last_status: dt.datetime | None = None

    print("[raw04] Шаг 2 из 3: рассчитываются новые таблицы справки")
    for record_number, record in enumerate(records.values(), start=1):
        if record["updated"] is None:
            diagnostics["without_update_date"] += 1
        if record["status_time"] is None:
            diagnostics["without_status_time"] += 1
        elif last_status is None or record["status_time"] > last_status:
            last_status = record["status_time"]
        if record["transport_type"] is None:
            diagnostics["unknown_transportation_type"] += 1
            if len(examples["unknown_transportation_type"]) < 20:
                examples["unknown_transportation_type"].append(record["shipment"])
        if record["transport"] is None:
            diagnostics["unknown_transport"] += 1
            if len(examples["unknown_transport"]) < 20:
                examples["unknown_transport"].append(record["shipment"])
        if record["country_from"] is None or record["country_to"] is None:
            diagnostics["unknown_country"] += 1
            if len(examples["unknown_country"]) < 20:
                examples["unknown_country"].append(record["shipment"])

        applied = _application_date(record, diagnostics)
        applied_in_period = applied is not None and applied >= PERIOD_START
        country_from, country_to = record["country_from"], record["country_to"]
        transport_type, transport = record["transport_type"], record["transport"]
        post = _match_point((record["point_raw"],), aliases)

        if applied_in_period and record["status"] in COUNTED_APPLICATION_STATUSES and country_from == "KZ" and country_to in COUNTRIES and transport_type == "transit" and post:
            group = "auto" if groups.get(post) == {"auto"} else "rail"
            transit[(group, post, country_to)] += 1

        if applied_in_period and record["status"] in COUNTED_APPLICATION_STATUSES and country_from == "KZ" and country_to in COUNTRIES and transport in {"auto", "rail"}:
            if transport_type == "export":
                export[(transport, country_to)] += 1
            elif transport_type == "trade":
                mutual[(transport, country_to)] += 1

        direction = TRACKING_DIRECTION_BY_COUNTRIES.get((country_from, country_to))
        tracking_confirmed = "казахстан" in record["intermediate"] or "рк" in record["intermediate"] or direction is not None
        if applied_in_period and record["status"] in COUNTED_APPLICATION_STATUSES and direction and tracking_confirmed and transport in {"auto", "rail"}:
            tracking[(transport, direction)] += 1

        if record["status"] == FINISHED_STATUS and record["status_time"] is not None and record["status_time"] >= PERIOD_START and transport in {"auto", "rail"}:
            if country_from == "KZ" and country_to in COUNTRIES:
                completed[(transport, "out", country_to)] += 1
            if country_to == "KZ" and country_from in COUNTRIES:
                completed[(transport, "in", country_from)] += 1
                completed_in_all[country_from] += 1

        if (applied_in_period and record["status"] in COUNTED_APPLICATION_STATUSES
                and country_from == "KZ" and country_to in {"RU", "BY", "KG"}
                and transport_type in {"transit", "export", "trade"}):
            for number in record["np_numbers"]:
                applications[(country_to, transport_type)] += 1
                if post is None or post not in groups:
                    unknown_post_pairs += 1
                else:
                    group = "auto" if groups.get(post) == {"auto"} else "rail"
                    post_applications[(group, post, country_to)] += 1

        if applied_in_period and country_from == "KZ" and record["np_numbers"] and post is None:
            diagnostics["unknown_start_point"] += 1
            if len(examples["unknown_start_point"]) < 20:
                examples["unknown_start_point"].append(f"{record['shipment']}: {record['point_raw']}")
        if record_number % 25_000 == 0:
            print("[raw04] обработано актуальных перевозок: ", record_number, sep="")

    auto_t1 = _post_rows(transit, auto_posts, "auto")
    rail_t1 = _post_rows(transit, rail_posts, "rail")
    table1 = {"table": "Таблица1", "daily_raw04": True, "transport_split": True, "auto_rows": auto_t1, "rail_rows": rail_t1, "rail_registered_total": ""}
    table2 = {"table": "Таблица2", "daily_raw04": True, "transport_split": True, "export_split_rows": _split_country_rows(export), "trade_split_rows": _split_country_rows(mutual)}
    directions = [label for _, _, label in [("BY", "KG", "РБ - КР"), ("RU", "KG", "РФ - КР"), ("KG", "RU", "КР - РФ"), ("KG", "BY", "КР - РБ")]]
    tracking_rows = []
    for group, label in (("auto", "Авто"), ("rail", "ЖД")):
        row = {"Тип транспорта": label}
        for direction_name in directions: row[direction_name] = int(tracking[(group, direction_name)])
        row["Итого"] = sum(row[name] for name in directions)
        tracking_rows.append(row)
    tracking_rows.append({"Тип транспорта": "Итого", **{name: sum(row[name] for row in tracking_rows) for name in directions}, "Итого": sum(row["Итого"] for row in tracking_rows)})
    table3 = {"table": "Таблица3", "daily_raw04": True, "rows": tracking_rows, "directions": directions}
    table4 = {"table": "Таблица4", "daily_raw04": True, "auto_out_rows": _flow_rows(completed, "auto", False), "auto_in_rows": _flow_rows(completed, "auto", True), "rail_out_rows": _flow_rows(completed, "rail", False), "rail_in_rows": _flow_rows(completed, "rail", True)}
    table5 = {"table": "Таблица5", "daily_raw04": True, "rows": _country_rows(completed_in_all)}
    direction_rows = []
    for code, label in (("RU", "РК - РФ"), ("BY", "РК - РБ"), ("KG", "РК - КР")):
        row = {"Направление": label, "Транзит": int(applications[(code, "transit")]), "Экспорт": int(applications[(code, "export")]), "Взаимная торговля": int(applications[(code, "trade")])}
        row["Итого"] = row["Транзит"] + row["Экспорт"] + row["Взаимная торговля"]
        direction_rows.append(row)
    direction_rows.append({"Направление": "РК - РК (пилот)", **PILOT_DIRECTION_COUNTS, "Итого": PILOT_DIRECTION_TOTAL})
    direction_rows.append({"Направление": "Итого", **{name: sum(row[name] for row in direction_rows) for name in ("Транзит", "Экспорт", "Взаимная торговля", "Итого")}})
    table6 = {"table": "Таблица6", "daily_raw04": True, "rows": direction_rows, "pilot_static": PILOT_DIRECTION_TOTAL}
    auto_t7, rail_t7 = _post_rows(post_applications, auto_posts, "auto"), _post_rows(post_applications, rail_posts, "rail")
    table7 = {"table": "Таблица7", "daily_raw04": True, "auto_rows": auto_t7, "rail_rows": rail_t7, "auto_total": auto_t7[-1]["Итого"], "rail_total": rail_t7[-1]["Итого"], "total": sum(applications.values()) - unknown_post_pairs, "unresolved_points": unknown_post_pairs}
    diagnostics_rows = [
        {"Показатель": "Исходных строк", "Значение": source_count},
        {"Показатель": "Актуальных перевозок", "Значение": len(records)},
        {"Показатель": "Перевозок без номера", "Значение": missing_shipment},
        {"Показатель": "Перевозок без даты обновления", "Значение": diagnostics["without_update_date"]},
        {"Показатель": "Перевозок без времени статуса", "Значение": diagnostics["without_status_time"]},
        {"Показатель": "Замена даты активации датой создания", "Значение": diagnostics["activation_date_replaced_by_created"]},
        {"Показатель": "НП вне справочника пунктов", "Значение": unknown_post_pairs},
        {"Показатель": "Итог Таблицы6", "Значение": direction_rows[-1]["Итого"]},
        {"Показатель": "Пилот РК - РК (статично, только в Таблице6)", "Значение": PILOT_DIRECTION_TOTAL},
        {"Показатель": "Итог Таблицы7 + НП вне справочника + пилот", "Значение": table7["total"] + unknown_post_pairs + PILOT_DIRECTION_TOTAL},
    ]
    for kind in ("unknown_transportation_type", "unknown_transport", "unknown_country", "unknown_start_point"):
        diagnostics_rows.append({"Показатель": f"{kind}: {', '.join(examples[kind]) or '-'}", "Значение": diagnostics[kind]})
    diagnostic = {"table": "Диагностика", "daily_raw04": True, "rows": diagnostics_rows, "period_start": PERIOD_START.isoformat(), "last_status": last_status.isoformat() if last_status else "-"}
    print(
        "[raw04] Шаг 2 завершен: НП по направлениям=", direction_rows[-1]["Итого"],
        ", НП по распознанным постам=", table7["total"],
        ", вне справочника=", unknown_post_pairs,
        sep="",
    )
    return [table1, table2, table3, table4, table5, table6, table7, diagnostic]
