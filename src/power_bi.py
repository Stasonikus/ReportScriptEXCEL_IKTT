from __future__ import annotations

import datetime as dt
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


SHEET_NAMES = [
    "01_Перевозки",
    "02_Пломбы",
    "03_Страны",
    "04_Места (3)",
    "Количество перевозок",
]

TRANSPORT_HEADERS = [
    "ID", "Дата_среза", "Этап", "Показатель", "Направление",
    "Транспорт", "Пункт_пропуска_место", "Код_страны", "Страна",
    "Количество_перевозок", "Источник",
]

SEAL_HEADERS = [
    "ID", "Дата_среза", "Место", "Тип_места", "Остаток_НП",
    "Использовано_НП_всего", "В_пути",
    "Использовано_НП_текущий_месяц", "Примечание",
]

COUNTRY_ROWS = [
    ["RF", "Российская Федерация", "РФ"],
    ["RB", "Республика Беларусь", "РБ"],
    ["KG", "Кыргызская Республика", "КР"],
    ["AM", "Республика Армения", "РА"],
]

COUNTRIES = {
    "РФ": ("RF", "Российская Федерация"),
    "РБ": ("RB", "Республика Беларусь"),
    "КР": ("KG", "Кыргызская Республика"),
    "РА": ("AM", "Республика Армения"),
}

# Справочник повторяет лист "04_Места (3)" эталонного файла пользователя.
# Порядок строк также является частью формата.
LOCATION_ROWS = [
    ["тп «Аксу»", 52.115231, 76.874921, None, "Железнодорожный пункт пропуска", "Аксу"],
    ["тп «Алимбет»", 51.080005, 58.474443, None, "Автомобильный пункт пропуска", "Алимбет"],
    ["тп «Алтынколь»", 44.165030, 80.293626, None, "Железнодорожный пункт пропуска", "Алтынколь"],
    ["тп «Атамекен»", 40.836760, 68.560450, None, "Автомобильный пункт пропуска", "Атамекен"],
    ["тп «Ауыл»", 51.208258, 81.115648, None, "Автомобильный пункт пропуска", "Ауыл"],
    ["тп «Б. Конысбаева»", 40.971312, 68.728012, None, "Автомобильный пункт пропуска", "Б. Конысбаева"],
    ["тп «Бахты»", 46.678705, 82.777341, None, "Автомобильный пункт пропуска", "Бахты"],
    ["тп «Болашак»", 42.330853, 53.676747, None, "Железнодорожный пункт пропуска", "Болашак"],
    ["тп «Достык»", 45.233267, 82.520671, None, "Автомобильный пункт пропуска", "Достык"],
    ["тп «Жайсан»", 50.856127, 56.147735, None, "Автомобильный пункт пропуска", "Жайсан"],
    ["тп «Жайык»", 51.192850, 51.872659, None, "Железнодорожный пункт пропуска", "Жайык"],
    ["тп «Жана-Жол»", 55.094160, 68.275707, None, "Автомобильный пункт пропуска", "Жана-Жол"],
    ["тп «Казыгурт»", 41.524998, 69.406126, None, "Автомобильный пункт пропуска", "Казыгурт"],
    ["тп «Кайрак»", 54.000600, 61.638811, None, "Автомобильный пункт пропуска", "Кайрак"],
    ["тп «Капланбек»", 41.456281, 69.206840, None, "Автомобильный пункт пропуска", "Капланбек"],
    ["тп «Карасу»", 43.000666, 74.866559, None, "Автомобильный пункт пропуска", "Карасу"],
    ["тп «Кордай»", 43.025727, 74.706225, None, "Автомобильный пункт пропуска", "Кордай"],
    ["тп «Косак»", 53.398310, 77.704045, None, "Автомобильный пункт пропуска", "Косак"],
    ["тп «Курмангазы»", 46.553251, 48.763838, None, "Автомобильный пункт пропуска", "Курмангазы"],
    ["тп «Майкапчагай»", 47.493650, 85.608020, None, "Автомобильный пункт пропуска", "Майкапчагай"],
    ["тп «Нур Жолы»", 44.157365, 80.398653, None, "Автомобильный пункт пропуска", "Нур Жолы"],
    ["тп «Сарыагаш»", 41.465066, 69.147490, None, "Железнодорожный пункт пропуска", "Сарыагаш"],
    ["тп «Сырым»", 51.666034, 51.005950, None, "Автомобильный пункт пропуска", "Сырым"],
    ["тп «Тажен»", 44.896495, 55.996342, None, "Автомобильный пункт пропуска", "Тажен"],
    ["тп «Таскала»", 51.199654, 49.910863, None, "Автомобильный пункт пропуска", "Таскала"],
    ["тп «Темир баба»", 41.924233, 52.661838, None, "Автомобильный пункт пропуска", "Темир баба"],
    ["тп «Урлитобе»", 53.794263, 75.052727, None, "Автомобильный пункт пропуска", "Урлитобе"],
    ["тп «Шынгырлау»", 51.086243, 54.080037, None, "Железнодорожный пункт пропуска", "Шынгырлау"],
]

LOCATION_HEADERS = [
    "Место", "Широта", "Долгота", "Регион", "Тип", "Комментарий",
    "Пункт_перевозки",
]

SEAL_NOTE = (
    "Остаток — пломбы со статусом «На складе» в текущем "
    "местонахождении пункта."
)


def _norm(value: Any) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    for old, new in (
        ("ө", "о"), ("ұ", "у"), ("ү", "у"), ("қ", "к"),
        ("ғ", "г"), ("ә", "а"), ("і", "и"), ("ң", "н"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\b(тп|т п|ст)\b", " ", text)
    text = text.replace("железнодорожный", " ").replace("автомобильный", " ")
    text = text.replace("ж д", " ").replace("жд", " ").replace("авто", " ")
    return re.sub(r"\s+", " ", re.sub(r"[^0-9a-zа-я]+", " ", text)).strip()


def _as_number(value: Any) -> int:
    if value in (None, "", "-"):
        return 0
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _is_total(value: Any) -> bool:
    return _norm(value).startswith("итого")


def _country_from_title(value: Any) -> tuple[str, str]:
    text = _norm(value)
    if "беларус" in text:
        return "RB", "Республика Беларусь"
    if "кыргыз" in text or "киргиз" in text:
        return "KG", "Кыргызская Республика"
    if "армени" in text:
        return "AM", "Республика Армения"
    if "росси" in text or text == "рф":
        return "RF", "Российская Федерация"
    return "", "Не определена"


def _stage_name(table: Dict[str, Any], method: str) -> str:
    if table.get("stage") == "full-period" or method.startswith("5 -"):
        return "Единый период с 11.02.2026"
    if table.get("stage") == "stage2-direct" or method.startswith(("2 -", "3 -")):
        return "Второй этап"
    return "Первый этап"


def _source_for(table: Dict[str, Any], transport: str, method: str) -> str:
    if method.startswith(("3 -", "5 -")):
        return "Большая выгрузка"
    if transport == "ЖД" and table.get("stage") != "stage2-direct":
        return "ЖД выгрузка"
    return "Большая выгрузка" if table.get("stage") == "stage2-direct" else "report"


def _append_transport_row(
    rows: List[List[Any]], snapshot_date: dt.date, stage: str,
    indicator: str, direction: str, transport: str, point: str,
    country_code: str, country: str, value: Any, source: str,
) -> None:
    numeric = _as_number(value)
    if numeric == 0:
        return
    rows.append([
        len(rows) + 1, snapshot_date, stage, indicator, direction, transport,
        point or "Не определён", country_code or None,
        country or "Не определена", numeric, source,
    ])


def _add_transit_rows(
    rows: List[List[Any]], table: Dict[str, Any], snapshot_date: dt.date,
    method: str,
) -> None:
    if table.get("transport_split"):
        sections = [
            ("Авто", table.get("auto_rows", [])),
            ("ЖД", table.get("rail_rows", [])),
        ]
    else:
        sections = [("Не определён", table.get("rows", []))]

    for transport, source_rows in sections:
        for item in source_rows:
            point = str(item.get("ТП/ЖД") or "").strip()
            if not point or _is_total(point):
                continue
            for column, (country_code, country) in COUNTRIES.items():
                _append_transport_row(
                    rows, snapshot_date, _stage_name(table, method),
                    "Транзит (наложение НП)", "Транзит", transport, point,
                    country_code, country, item.get(column),
                    _source_for(table, transport, method),
                )


def _add_export_trade_rows(
    rows: List[List[Any]], table: Dict[str, Any], snapshot_date: dt.date,
    method: str,
) -> None:
    for indicator, key in (("Экспорт", "export"), ("Взаимная торговля", "trade")):
        if table.get("transport_split"):
            for item in table.get(f"{key}_split_rows", [])[:4]:
                country_code, country = _country_from_title(item.get("Страна"))
                for transport in ("Авто", "ЖД"):
                    _append_transport_row(
                        rows, snapshot_date, _stage_name(table, method),
                        indicator, "Из РК", transport, "Не определён",
                        country_code, country, item.get(transport),
                        _source_for(table, transport, method),
                    )
        else:
            for item in table.get(f"{key}_rows", [])[:4]:
                country_code, country = _country_from_title(item.get("Страна"))
                _append_transport_row(
                    rows, snapshot_date, _stage_name(table, method),
                    indicator, "Из РК", "Не определён", "Не определён",
                    country_code, country, item.get("Количество перевозок"),
                    _source_for(table, "", method),
                )


def _add_tracking_rows(
    rows: List[List[Any]], table: Dict[str, Any], snapshot_date: dt.date,
    method: str,
) -> None:
    destination_countries = {
        "РБ - КР": ("KG", "Кыргызская Республика"),
        "РФ - КР": ("KG", "Кыргызская Республика"),
        "КР - РФ": ("RF", "Российская Федерация"),
        "КР - РБ": ("RB", "Республика Беларусь"),
    }
    for item in table.get("rows", []) or []:
        transport = str(item.get("Тип транспорта") or "").strip()
        if transport not in {"Авто", "ЖД"}:
            continue
        for direction, (country_code, country) in destination_countries.items():
            _append_transport_row(
                rows, snapshot_date, _stage_name(table, method),
                "Услуги прослеживания", direction, transport,
                "Не определён", country_code, country, item.get(direction),
                "Большая выгрузка",
            )


def _add_seal_application_rows(
    rows: List[List[Any]], table: Dict[str, Any], snapshot_date: dt.date,
    method: str,
) -> None:
    """Add seal-application facts by normalized customs post."""
    sections = (("Авто", table.get("auto_rows", [])), ("ЖД", table.get("rail_rows", [])))
    for transport, source_rows in sections:
        for item in source_rows:
            point = str(item.get("ТП/ЖД") or "").strip()
            if not point or _is_total(point):
                continue
            for column, (country_code, country) in COUNTRIES.items():
                _append_transport_row(
                    rows, snapshot_date, _stage_name(table, method),
                    "Наложение НП", "Из РК", transport, point,
                    country_code, country, item.get(column),
                    "Большая выгрузка",
                )


def build_power_bi_facts(
    tables: Iterable[Dict[str, Any]], report_date: dt.date, method: str,
    point_registry_path: Path | None = None,
) -> tuple[List[List[Any]], List[List[Any]]]:
    """Create data blocks matching the supplied Power BI workbook schema."""
    del point_registry_path  # retained for compatibility with existing callers
    table_list = list(tables)
    snapshot_date = next(
        (
            table.get("date")
            for table in table_list
            if table.get("table") == "Таблица5"
            and isinstance(table.get("date"), dt.date)
        ),
        report_date,
    )

    transport_rows: List[List[Any]] = []
    for table in table_list:
        if table.get("table") == "Таблица1":
            _add_transit_rows(transport_rows, table, snapshot_date, method)
        elif table.get("table") == "Таблица2":
            _add_export_trade_rows(transport_rows, table, snapshot_date, method)
        elif table.get("table") == "Таблица3":
            _add_tracking_rows(transport_rows, table, snapshot_date, method)
        elif table.get("table") == "Таблица7":
            _add_seal_application_rows(transport_rows, table, snapshot_date, method)

    return transport_rows, _build_seal_rows(table_list, snapshot_date)


def _indexed_values(
    rows: Iterable[Dict[str, Any]], name_key: str, value_key: str,
) -> Dict[str, int]:
    result: Dict[str, int] = {}
    for row in rows or []:
        name = str(row.get(name_key) or "").strip()
        if not name or _is_total(name):
            continue
        result[_norm(name)] = _as_number(row.get(value_key))
    return result


def _lookup(index: Dict[str, int], *names: str) -> int:
    for name in names:
        key = _norm(name)
        if key in index:
            return index[key]
        for existing, value in index.items():
            if key and existing and (key in existing or existing in key):
                return value
    return 0


def _build_seal_rows(
    tables: Iterable[Dict[str, Any]], snapshot_date: dt.date,
) -> List[List[Any]]:
    table5 = next((table for table in tables if table.get("table") == "Таблица5"), {})
    stock = _indexed_values(table5.get("rows", []), "Точка", "Количество НП")
    used = _indexed_values(table5.get("used_np_rows", []), "Пункт", "Количество")
    monthly = _indexed_values(
        table5.get("monthly_used_np_rows", table5.get("used_np_rows", [])),
        "Пункт", "Количество",
    )

    aliases = {
        "Алтынколь": ("Алтынколь-жол",),
        "Сарыагаш": ("ст. Сарыагаш",),
        "Темир баба": ("Темир-Баба",),
        "Нур Жолы": ("Нур жолы",),
    }
    output: List[List[Any]] = []
    for row_id, location in enumerate(LOCATION_ROWS, start=1):
        place, _, _, _, _, point_name = location
        search_names = (point_name, *aliases.get(point_name, ()))
        output.append([
            row_id, snapshot_date, place, "Пункт пропуска",
            _lookup(stock, *search_names), _lookup(used, *search_names), 0,
            _lookup(monthly, *search_names), SEAL_NOTE,
        ])
    return output


def _write_plain_sheet(
    workbook: Workbook, title: str, headers: List[str],
    rows: Iterable[Iterable[Any]], widths: List[float],
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
    )
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def write_power_bi_output(
    out_path: Path, tables: Iterable[Dict[str, Any]], report_date: dt.date,
    method: str, point_registry_path: Path | None = None,
) -> None:
    table_list = list(tables)
    transport_rows, seal_rows = build_power_bi_facts(
        table_list, report_date, method, point_registry_path,
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_plain_sheet(
        workbook, "01_Перевозки", TRANSPORT_HEADERS, transport_rows,
        [6, 21, 13, 24, 13, 14, 22, 12, 23, 22, 45],
    )
    _write_plain_sheet(
        workbook, "02_Пломбы", SEAL_HEADERS, seal_rows,
        [4, 21, 20, 16, 12, 23, 8, 31, 45],
    )
    _write_plain_sheet(
        workbook, "03_Страны", ["Column1", "Column2", "Column3"],
        COUNTRY_ROWS, [9, 23, 9],
    )

    location_comment = "Пункт использован в перевозках с 11.02.2026"
    locations = [row[:5] + [location_comment, row[5]] for row in LOCATION_ROWS]
    _write_plain_sheet(
        workbook, "04_Места (3)", LOCATION_HEADERS, locations,
        [20, 11, 11, 8, 32, 45, 17],
    )
    _write_plain_sheet(
        workbook, "Количество перевозок",
        ["Column", "Количество перевозок"],
        [["Всего", sum(row[9] for row in transport_rows)]], [8, 22],
    )

    for sheet_name in ("01_Перевозки", "02_Пломбы"):
        worksheet = workbook[sheet_name]
        for cell in worksheet["B"][1:]:
            cell.number_format = "YYYY-MM-DD HH:MM:SS"

    workbook.save(out_path)
    print(
        f"[power_bi] wrote {len(transport_rows)} transport rows and "
        f"{len(seal_rows)} seal rows to: {out_path}"
    )
