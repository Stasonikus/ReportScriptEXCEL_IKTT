from __future__ import annotations

import datetime as dt
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Описание работы daily_script.docx"

FONT = "Times New Roman"
FONT_SIZE = Pt(14)
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF2F8"
GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, enabled: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = enabled


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char1, instr_text, fld_char2))


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = FONT_SIZE
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    for name, size, before, after in (
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 14, 9, 4),
        ("Heading 3", 14, 7, 3),
    ):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Document note" not in styles:
        note = styles.add_style("Document note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = styles["Document note"]
    note.font.name = FONT
    note.font.size = Pt(12)
    note.font.italic = True
    note.font.color.rgb = RGBColor(70, 70, 70)
    note._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_after = Pt(5)

    if "Table text" not in styles:
        table_style = styles.add_style("Table text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_style = styles["Table text"]
    table_style.font.name = FONT
    table_style.font.size = Pt(10.5)
    table_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    table_style.paragraph_format.first_line_indent = Cm(0)
    table_style.paragraph_format.space_before = Pt(0)
    table_style.paragraph_format.space_after = Pt(0)
    table_style.paragraph_format.line_spacing = 1.0


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(4.5)
    r = p.add_run("СЛУЖЕБНОЕ ОПИСАНИЕ")
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(1.0)
    p.paragraph_format.space_after = Cm(1.0)
    r = p.add_run(
        "программы daily_script для автоматического формирования\n"
        "аналитической отчетности по навигационным пломбам"
    )
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(18)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    values = (
        ("Версия программы", "2.4.2"),
        ("Дата актуализации описания", "21.09.2026"),
        ("Начало второго этапа", "31.07.2026 00:00"),
        ("Основной результат", "Excel-отчет и набор данных для Power BI"),
    )
    for row, values_row in zip(table.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles["Table text"]
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], LIGHT_BLUE)

    p = doc.add_paragraph(style="Document note")
    p.paragraph_format.space_before = Cm(1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Документ описывает фактическую логику сборки отчета, доступные способы "
        "расчета, требования к входным файлам и правила формирования показателей."
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Cm(5.0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Республика Казахстан, 2026").italic = True
    p.add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, *, note: bool = False, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="Document note" if note else "Normal")
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.style = doc.styles["Normal"]
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.hanging_indent = Cm(0.45)
        p.add_run("• ")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.hanging_indent = Cm(0.55)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_together = True
        p.add_run(f"{number}. ").bold = True
        p.add_run(item)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.style = doc.styles["Table text"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = str(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles["Table text"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    footer = section.footer.paragraphs[0]
    footer.style = doc.styles["Table text"]
    add_page_number(footer)

    add_title_page(doc)

    add_heading(doc, "1. Общие сведения о программе")
    add_para(
        doc,
        "daily_script версии 2.4.2 предназначен для автоматизированного "
        "формирования ежедневной аналитической отчетности по навигационным "
        "пломбам и перевозкам. Программа считывает входные Excel-файлы, "
        "нормализует названия стран, пунктов пропуска, статусы и виды транспорта, "
        "выполняет расчеты и формирует два согласованных результата: служебный "
        "пятилистовой отчет и плоский набор данных для Power BI."
    )
    add_para(
        doc,
        "Программа поставляется как консольный EXE-файл. Для обычного запуска "
        "пользователю не требуются Python, среда разработки или ручное изменение "
        "формул. Выбор логики выполняется в меню способов расчета."
    )
    add_para(
        doc,
        "Все расчеты второго этапа используют контрольную дату 31.07.2026 00:00. "
        "Для большой выгрузки попадание записи во второй этап определяется по "
        "столбцу «Дата создания». Записи с более ранней датой не включаются в "
        "расчеты второго этапа."
    )

    add_heading(doc, "2. Рабочие папки и состав программы")
    add_para(
        doc,
        "При запуске программа использует папки, расположенные рядом с EXE-файлом. "
        "Если необходимые каталоги отсутствуют, они создаются автоматически."
    )
    add_table(
        doc,
        ["Папка или файл", "Назначение"],
        [
            ["in", "Входные файлы текущего запуска."],
            ["in/processed", "Архив успешно использованных входных файлов с датой и временем обработки в имени."],
            ["out/results", "Служебные отчеты Excel с листами Таблица1–Таблица5."],
            ["out/power bi data", "Нормализованные Excel-файлы для загрузки в Power BI."],
            ["monthly_baselines", "Снимки накопительного расхода НП по датам для расчета показателя текущего месяца."],
            ["data", "Встроенные справочники пунктов, форматов номеров ТС и базы второго этапа; включаются в EXE."],
        ],
        [4.2, 11.8],
    )
    add_para(
        doc,
        "Исходные файлы перемещаются в in/processed только после успешной записи "
        "обоих выходных файлов. Повторный файл не перезаписывает архивный: к имени "
        "добавляется временная метка, а при совпадении — дополнительный счетчик."
    )

    add_heading(doc, "3. Входные файлы")
    add_heading(doc, "3.1. Обычный статистический отчет report_*.xlsx", level=2)
    add_para(
        doc,
        "Файл должен иметь имя report_*.xlsx и лист «Отчёт». Это основной источник "
        "старого метода, источник Таблицы 4 в Основном режиме (внутренний код 3) "
        "и источник старой части смешанного режима 4. Если файл по структуре "
        "распознан как большая выгрузка, "
        "он не принимается одновременно как обычный report."
    )

    add_heading(doc, "3.2. Большая выгрузка перевозок", level=2)
    add_para(
        doc,
        "Большая выгрузка определяется по структуре заголовков, а не по имени. "
        "Поддерживаются настоящий XLSX и HTML-таблица, сохраненная системой с "
        "расширением XLS. Строка заголовков может находиться в первых 50 строках."
    )
    add_table(
        doc,
        ["Группа", "Обязательные поля"],
        [
            ["Идентификация и даты", "Дата создания; Номер перевозки; Статус; Дата обновления"],
            ["Классификация", "Тип перевозки; Тип транспорта; Владелец; Компания; БИН"],
            ["ТС и пломбы", "Номера ТС; Номера НП; Номера документов; Коды ТН ВЭД"],
            ["Маршрут", "Начальный пункт; Страна начала; Конечный пункт; Страна завершения; Промежуточные точки; Промежуточные страны"],
            ["Счетчики", "Кол. путевых точек; Кол. водителей; Кол. операторов"],
        ],
        [4.2, 11.8],
    )
    add_para(
        doc,
        "Версия 2.4.2 принимает переименованные заголовки «Водителей», «Путевых "
        "точек», «Операторов» и «Промежуточные пункты», автоматически приводя их "
        "к каноническим полям расчета."
    )

    add_heading(doc, "3.3. Файл готовности АПП/ЖД/Территория", level=2)
    add_para(
        doc,
        "Имя файла должно начинаться со слова «Готовность» и содержать признаки "
        "АПП и Территория. Поддерживается XLSX. Программа выбирает лист, в названии "
        "которого присутствуют слова «распределение» и «НП»; если такого листа нет, "
        "используется первый лист. Файл содержит текущие остатки и распределение НП "
        "по пунктам, складам и территориям."
    )

    add_heading(doc, "3.4. Железнодорожная выгрузка", level=2)
    add_para(
        doc,
        "Поддерживаются файлы вида «ЖД Выгрузка*.xlsx», бинарные XLS и HTML-XLS. "
        "Источник применяется в дополнительных режимах 1 и 4 для уточнения железнодорожных "
        "показателей. Для XLS программа находит строку с полями «Владелец записи», "
        "«Статус», company_id, start waypoint и end waypoint."
    )
    add_para(
        doc,
        "В папке in должен находиться не более чем один файл каждого требуемого "
        "типа. Временные файлы Excel с префиксом ~$ игнорируются.",
        note=True,
    )

    add_heading(doc, "4. Выбор способа расчета")
    add_para(
        doc,
        "После запуска отображается главное меню из трех пунктов. Наиболее часто "
        "используемый прежний режим 3 вынесен в отдельный пункт «1 — Основной "
        "режим». Его расчетная логика не изменена."
    )
    add_table(
        doc,
        ["Пункт", "Наименование", "Действие"],
        [
            ["1", "Основной режим", "Сразу запускает прежний режим 3: большая выгрузка с разделением авто/ЖД, report и файл готовности."],
            ["2", "Остальные режимы", "Открывает дополнительное меню со старым, новым и смешанным способами расчета."],
            ["3", "Выход", "Завершает программу без обработки файлов."],
        ],
        [1.4, 4.0, 10.6],
    )
    add_para(
        doc,
        "При выборе пункта 2 открывается дополнительное меню. Пункт 0 возвращает "
        "пользователя в главное меню и не запускает обработку."
    )
    add_table(
        doc,
        ["Пункт", "Режим", "Необходимые входы", "Источник таблиц"],
        [
            ["0", "Назад", "Нет", "Возврат в главное меню."],
            ["1", "Старый метод", "report; готовность; ЖД выгрузка", "Т1, Т2, Т3 и Т5 — report; ЖД-уточнение — ЖД выгрузка; Т4 — report + готовность."],
            ["2", "Новый метод", "большая выгрузка; готовность", "Т1–Т5 рассчитываются из большой выгрузки; готовность задает текущие остатки Т4."],
            ["4", "Смешанный метод", "большая выгрузка; report; готовность; ЖД выгрузка", "Т1–Т2 — большая выгрузка с подтвержденными ЖД-данными; Т3–Т5 — старый метод."],
        ],
        [1.0, 3.0, 4.2, 7.8],
    )
    add_para(
        doc,
        "В дополнительном режиме 1 отсутствие необязательного файла может быть "
        "подтверждено пользователем, однако соответствующая часть отчета будет "
        "неполной. Для Основного режима report и файл готовности обязательны; для "
        "смешанного режима 4 также обязательна ЖД-выгрузка."
    )

    add_heading(doc, "5. Общие правила обработки большой выгрузки")
    add_numbered(
        doc,
        [
            "Проверяется полный набор обязательных заголовков. Файл с неподходящей структурой не принимается.",
            "Дата создания преобразуется в дату и время. В расчеты второго этапа попадают только записи с датой не ранее 31.07.2026 00:00.",
            "Записи группируются по номеру перевозки. Если номер повторяется, сохраняется строка с наиболее поздней датой обновления; при равенстве дат — последняя строка файла.",
            "Номера НП из всех строк одного номера перевозки объединяются без потери уникальных значений.",
            "Страны приводятся к кодам KZ, RU, BY, KG и ARM. В целевых таблицах используются Россия, Беларусь, Кыргызстан и Армения.",
            "Статусы, виды перевозок, названия пунктов и номера ТС нормализуются перед подсчетом.",
        ],
    )
    add_para(
        doc,
        "Пустая или некорректная «Дата создания» не заменяется текущей датой. Такая "
        "строка остается в диагностической статистике, но не участвует в расчетах "
        "второго этапа."
    )

    add_heading(doc, "6. Определение вида транспорта и пункта пропуска")
    add_heading(doc, "6.1. Автомобильный или железнодорожный транспорт", level=2)
    add_para(
        doc,
        "В режимах с большой выгрузкой вид транспорта определяется в следующем "
        "порядке: сначала по формату значений «Номера ТС» из встроенного справочника, "
        "затем по подтверждению номером перевозки из ЖД-выгрузки (только дополнительный режим 4), "
        "после этого по столбцу «Тип транспорта». Если классификация отсутствует, "
        "для совместимости применяется автомобильная группа."
    )
    add_para(
        doc,
        "Для Таблицы 1 окончательная группа также контролируется справочником "
        "пунктов: пункт, зарегистрированный только как авто или только как ЖД, "
        "попадает в соответствующий раздел независимо от неоднозначного номера ТС."
    )

    add_heading(doc, "6.2. Сопоставление пунктов", level=2)
    add_para(
        doc,
        "Пункт определяется по нормализованному значению «Начальный пункт». "
        "Сравнение выполняется с каноническим названием, общими aliases и "
        "вариантами написания для большой и ЖД-выгрузки из "
        "border_crossing_points.json. Нормализация убирает регистр, кавычки, "
        "служебные знаки, лишние пробелы и унифицирует русские и казахские буквы."
    )
    add_table(
        doc,
        ["Автомобильные пункты Таблицы 1", "Железнодорожные/стыковочные пункты Таблицы 1"],
        [[
            "Алаколь; Казыгурт; Нур жолы; Калжат; Бахты; Тажен; Атамекен; Б. Конысбаева; Майкапчагай; Темир-баба; Капланбек",
            "Морпорт Актау; Курык порт; Болашак; Оазис; Сарыагаш; Алтынколь; Темиржол",
        ]],
        [8.0, 8.0],
    )
    add_para(
        doc,
        "Перевозка с нераспознанным начальным пунктом не распределяется по строкам "
        "Таблицы 1. Количество таких записей выводится в диагностике unresolved_t1."
    )

    add_heading(doc, "7. Правила формирования аналитических таблиц")
    add_heading(doc, "7.1. Таблица 1 — транзитные перевозки с наложением НП", level=2)
    add_para(
        doc,
        "Для большой выгрузки учитывается последняя версия уникальной перевозки при "
        "одновременном выполнении условий: дата создания не ранее 31.07.2026 00:00; "
        "статус «Активированная перевозка» или «Завершенная перевозка»; тип перевозки "
        "содержит признак транзита; страна начала — Казахстан; страна завершения — "
        "Россия, Беларусь, Кыргызстан или Армения; начальный пункт сопоставлен со "
        "справочником."
    )
    add_para(
        doc,
        "В Основном режиме и дополнительном режиме 4 таблица выводится двумя разделами: автомобильные и "
        "железнодорожные пункты. Столбцы РФ, РБ, КР и РА определяются по стране "
        "завершения. Показатель «Количество оформленных ЖД» отражает уникальные ЖД-"
        "перевозки второго этапа со статусом «Активированная» или «Завершенная», "
        "началом в РК и допустимым владельцем записи."
    )

    add_heading(doc, "7.2. Таблица 2 — экспорт и взаимная торговля из РК", level=2)
    add_para(
        doc,
        "Учитываются уникальные перевозки второго этапа со статусом «Активированная "
        "перевозка» или «Завершенная перевозка», началом в Казахстане и завершением "
        "в одной из четырех целевых стран. Тип перевозки должен быть распознан как "
        "«экспорт» либо «перевозка товаров в рамках взаимной торговли». В Основном "
        "режиме и дополнительном режиме 4 значения дополнительно разделяются на авто и ЖД."
    )

    add_heading(doc, "7.3. Таблица 3 — завершенные перевозки по направлениям", level=2)
    add_para(
        doc,
        "Таблица содержит четыре блока: авто из РК, авто в РК, ЖД из РК и ЖД в РК. "
        "Для направления из РК в большой выгрузке используется статус «Завершенная "
        "перевозка», допустимый владелец и страна назначения. Для направления в РК "
        "принимаются статусы «Завершенная перевозка» и «Деактивированная перевозка», "
        "страна назначения должна быть Казахстаном. В Основном режиме все четыре "
        "блока строятся из большой выгрузки; в дополнительном режиме 4 Таблица 3 "
        "берется из старого метода."
    )

    add_heading(doc, "7.4. Таблица 4 — остатки и расход навигационных пломб", level=2)
    add_para(
        doc,
        "Левая часть содержит текущий остаток НП по пунктам пропуска, складам и "
        "территориям из файла готовности. Накопительное количество использованных НП "
        "берется из источника соответствующего метода и для старого/report-расчета "
        "уменьшается на зафиксированную базу первого этапа на 31.07.2026. Остатки НП "
        "не уменьшаются на эту базу."
    )
    add_para(
        doc,
        "Показатель «Использовано НП в текущем месяце» рассчитывается отдельно: из "
        "текущего накопительного значения вычитается снимок на последний день "
        "предыдущего календарного месяца. При отсутствии точного снимка применяется "
        "последний доступный снимок предыдущего месяца с предупреждением; при полном "
        "отсутствии базы значение не рассчитывается. Средний расход в день делится "
        "на количество прошедших дней текущего месяца."
    )
    add_para(
        doc,
        "В Основном режиме Таблица 4 принципиально формируется только из report и файла "
        "готовности; сведения о НП большой выгрузки в нее не подмешиваются. После "
        "успешного запуска создается новый снимок в monthly_baselines."
    )

    add_heading(doc, "7.5. Таблица 5 — активные перевозки в сторону РК", level=2)
    add_para(
        doc,
        "Учитываются перевозки со статусом «Активированная перевозка», страной "
        "завершения Казахстан и страной начала Россия, Беларусь, Кыргызстан или "
        "Армения. В методах, использующих большую выгрузку, применяется дата создания "
        "не ранее 31.07.2026 00:00."
    )

    add_heading(doc, "8. Выходные файлы")
    add_para(
        doc,
        "За один успешный запуск формируются два файла с одинаковой временной меткой. "
        "Суффикс имени зависит от метода: без суффикса для метода 1, new method для "
        "метода 2, combined method для метода 3 и hybrid method для метода 4."
    )
    add_table(
        doc,
        ["Результат", "Расположение", "Содержание"],
        [
            ["Служебный отчет", "out/results/stage 2 report...xlsx", "Пять форматированных листов: Таблица1, Таблица2, Таблица3, Таблица4 и Таблица5."],
            ["Данные Power BI", "out/power bi data/stage 2 report...xlsx", "Плоские таблицы «Факты», «Пункты», «Страны» и «Параметры»."],
        ],
        [3.5, 5.0, 7.5],
    )
    add_para(
        doc,
        "Лист «Факты» содержит дату отчета, способ расчета, номер таблицы, показатель, "
        "транспорт, направление, код и название страны, пункт, вид границы, числовое "
        "значение и источник. Лист «Пункты» формируется из встроенного реестра из 54 "
        "справочных точек и содержит идентификатор, название, транспорт, тип границы "
        "и варианты написания. Нулевые факты в Power BI-файл не записываются."
    )
    add_para(
        doc,
        "Текущая версия реестра используется для классификации и подписей. Координаты "
        "карты в выходной файл версии 2.4.2 не записываются; для географической карты "
        "Power BI требуется дополнить справочник полями широты и долготы либо связать "
        "лист «Пункты» с отдельным координатным справочником по идентификатору пункта.",
        note=True,
    )

    add_heading(doc, "9. Порядок работы пользователя")
    add_numbered(
        doc,
        [
            "Закрыть входные и ранее сформированные Excel-файлы.",
            "Очистить папку in от файлов предыдущего незавершенного запуска; папку processed не удалять.",
            "Поместить в in ровно по одному файлу каждого типа, необходимого выбранному способу расчета.",
            "Запустить report_builder 2.4.2.exe.",
            "В главном меню выбрать 1 для Основного режима, 2 для открытия остальных режимов либо 3 для выхода.",
            "В дополнительном меню выбрать 1, 2 или 4; пункт 0 возвращает в главное меню. При необходимости подтвердить продолжение без необязательного файла.",
            "Дождаться сообщения [main] done. Не закрывать окно во время чтения и записи файлов.",
            "Проверить служебный отчет в out/results и набор для Power BI в out/power bi data.",
            "При необходимости сверить архив использованных источников в in/processed и месячный снимок в monthly_baselines.",
        ],
    )

    add_heading(doc, "10. Диагностика и нештатные ситуации")
    add_table(
        doc,
        ["Сообщение или ситуация", "Причина и действие пользователя"],
        [
            ["Файл источника не найден", "Нет подходящего report_*.xlsx либо выбран неверный метод. Проверить имя, лист «Отчёт» и папку in."],
            ["Большая выгрузка не найдена", "Нет XLSX/HTML-XLS с полным набором обязательных заголовков. Проверить строку заголовков и формат файла."],
            ["Найдено несколько файлов", "В in находится более одного источника одного типа. Оставить только файл текущего запуска."],
            ["Таблица 4 будет пропущена", "Файл готовности отсутствует, неоднозначен или не читается. Для методов 3 и 4 продолжение невозможно."],
            ["unresolved_t1 больше нуля", "Часть начальных пунктов не сопоставлена со справочником aliases. Проверить исходные названия и border_crossing_points.json."],
            ["invalid_creation_dates больше нуля", "Часть дат создания пуста или не распознана; такие записи не включены во второй этап."],
            ["Не удалось записать файл результата", "Закрыть открытый выходной Excel-файл, проверить доступ к папке out и свободное место."],
            ["База предыдущего месяца не найдена", "Месячный расход НП не рассчитывается. Добавить снимок последнего дня предыдущего месяца в monthly_baselines."],
        ],
        [5.0, 11.0],
    )
    add_para(
        doc,
        "Для технической сверки рекомендуется сохранять текст консоли: программа "
        "выводит количество прочитанных строк, уникальных перевозок, дублей, "
        "нераспознанных дат и пунктов, распределение авто/ЖД и итоговые значения "
        "расчетов."
    )

    add_heading(doc, "11. Контрольные правила версии 2.4.2")
    add_bullets(
        doc,
        [
            "Дата второго этапа всегда фильтруется по «Дате создания», а не по дате обновления.",
            "Повторный номер перевозки считается один раз по самой свежей версии записи.",
            "В Основном режиме (внутренний код 3) авто и ЖД для Таблиц 1, 2, 3 и 5 определяются без отдельной ЖД-выгрузки.",
            "Капланбек находится в автомобильном разделе после Темир-баба; Темиржол находится в железнодорожном разделе.",
            "Морпорт Актау и Курык порт относятся к железнодорожным/стыковочным пунктам отчета.",
            "Текущий месячный расход НП и накопительное использование второго этапа являются разными показателями и не должны совпадать автоматически.",
            "Служебный отчет и Power BI-файл должны быть сформированы в одном запуске и иметь одинаковую временную метку.",
        ],
    )

    add_para(
        doc,
        "Конец документа.",
        note=True,
    )
    return doc


def main() -> None:
    doc = build_document()
    doc.core_properties.title = "Описание работы daily_script 2.4.2"
    doc.core_properties.subject = "Автоматическое формирование отчетности по навигационным пломбам"
    doc.core_properties.author = "ReportScriptEXCEL_IKTT"
    doc.core_properties.comments = "Актуализировано по фактической логике версии 2.4.2"
    doc.core_properties.modified = dt.datetime(2026, 9, 21, 0, 0, 0)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
