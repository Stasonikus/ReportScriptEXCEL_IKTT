from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
VERSION = "2.7.0"
FONT = "Times New Roman"
BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
BORDER = "D9D9D9"


def font(run, size: float = 12, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{key}"), FONT)


def shade(cell, fill: str) -> None:
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(node)


def border(cell) -> None:
    holder = OxmlElement("w:tcBorders")
    for name in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), BORDER)
        holder.append(edge)
    cell._tc.get_or_add_tcPr().append(holder)


def cell_margins(cell) -> None:
    holder = OxmlElement("w:tcMar")
    for name in ("top", "start", "bottom", "end"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:w"), "100")
        edge.set(qn("w:type"), "dxa")
        holder.append(edge)
    cell._tc.get_or_add_tcPr().append(holder)


def setup(doc: Document, compact: bool = False) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11 if compact else 12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    title = doc.styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(19 if compact else 21)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    for name, size in (("Heading 1", 15), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    note = doc.styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
    note.font.name = FONT
    note.font.size = Pt(10.5)
    note.font.italic = True
    note._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    note.paragraph_format.space_after = Pt(6)


def add_text(doc: Document, text: str, note: bool = False) -> None:
    doc.add_paragraph(text, style="Note" if note else "Normal")


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.add_run("• ")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for number, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.85)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        lead = p.add_run(f"{number}. ")
        font(lead, 12, True)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        cell.width = Cm(widths[index])
        shade(cell, BLUE)
        border(cell)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                font(r, 9.5, True)
                r.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = value
            cell.width = Cm(widths[index])
            if row_index % 2:
                shade(cell, PALE_BLUE)
            border(cell)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    font(r, 9.3, index == 0)


def title(doc: Document, headline: str, subtitle: str) -> None:
    doc.add_paragraph(headline, style="Title")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    font(r, 13, True)


def technical() -> Document:
    doc = Document()
    setup(doc)
    title(doc, "Описание работы daily script", f"Версия {VERSION}")
    add_text(doc, "Документ описывает актуальную логику программы: меню, входные источники, фильтры и результат ежедневной справки RAW04.", True)

    doc.add_heading("1 Главное меню", level=1)
    add_table(doc, ["Пункт", "Режим", "Назначение"], [
        ["1", "Основной режим: Ежедневная справка RAW04", "Формирует ежедневную справку. Для запуска обязательна только большая выгрузка."],
        ["2", "Остальные режимы", "Открывает старый, новый, расчет с АПП КГД и большой выгрузкой, смешанный и единый режимы."],
        ["3", "Выход", "Завершает работу программы."],
    ], [1.4, 5.1, 9.0])
    add_text(doc, "В дополнительном меню пункт 3 называется «Режим v2.7.0: с АПП КГД и большой выгрузкой». Он сохраняет прежнюю логику расчета с отчетом КГД и большой выгрузкой.")

    doc.add_heading("2 Входные файлы основного режима RAW04", level=1)
    add_table(doc, ["Источник", "Статус", "Использование"], [
        ["Большая выгрузка", "Обязательный", "Основной источник перевозок. Определяется по структуре заголовков, а не по имени файла."],
        ["report_*.xlsx", "Необязательный", "Нужен для Таблицы8: остатки НП и месячный расход. При отсутствии расчет продолжается, Таблица8 помечается как недоступная."],
        ["Готовность АПП ЖД Территория.xlsx", "Необязательный", "Дополняет остатки НП. При отсутствии выводится предупреждение."],
    ], [3.7, 2.7, 9.1])
    add_text(doc, "Большая выгрузка должна содержать дату создания, номер перевозки, статус, дату обновления, тип перевозки, тип транспорта, страны начала и завершения, начальный и конечный пункты, номера НП и служебные поля. Вариант «Железнодорожный траспорт» распознается как ЖД.")

    doc.add_heading("3 Правила отбора большой выгрузки", level=1)
    add_bullets(doc, [
        "Период основных расчетов: с 11 февраля 2026 года по последнюю дату в источнике.",
        "Повторяющиеся номера перевозок объединяются: используется актуальная версия записи.",
        "Таблица1: таможенный транзит из РК, активированные и завершенные перевозки, распределение по постам, странам и авто/ЖД.",
        "Таблица2: экспорт и взаимная торговля из РК, активированные и завершенные перевозки, страны и авто/ЖД.",
        "Таблица3: прослеживание через РК по направлениям РБ-КР, РФ-КР, КР-РФ и КР-РБ; статусы активированная и завершенная перевозка.",
        "Таблица4: завершенные перевозки, отдельно из РК и в РК, с разбиением на авто и ЖД.",
        "Таблица5: активные перевозки в РК по направлениям.",
        "Таблица6: наложение НП по направлениям, включая статичную строку РК-РК (пилот).",
        "Таблица7: наложение НП по точкам начала и нормализованным названиям постов.",
    ])

    doc.add_heading("4 Листы результата RAW04", level=1)
    add_table(doc, ["Лист", "Содержание"], [
        ["Шапка", "Числа для текста справки: оформлено, мониторинг, завершено и разбивка по процедурам и странам."],
        ["Таблица1", "Транзит по пунктам, странам назначения и типам транспорта."],
        ["Таблица2", "Экспорт и взаимная торговля по странам и типам транспорта."],
        ["Таблица3", "Услуги прослеживания через РК."],
        ["Таблица4", "Завершенные перевозки из РК и в РК по авто и ЖД."],
        ["Таблица5", "Активные перевозки в сторону РК."],
        ["Таблица6", "Наложения НП по направлениям."],
        ["Таблица7", "Наложения НП по автомобильным постам и ЖД-станциям."],
        ["Таблица8", "Остатки и месячный расход НП из дополнительного report; при его отсутствии содержит понятное уведомление."],
        ["Диагностика", "Контроль нормализации постов и нераспознанных значений."],
    ], [3.0, 12.5])

    doc.add_heading("5 Рабочие папки", level=1)
    add_bullets(doc, [
        "in: файлы очередного запуска.",
        "in/processed: архив успешно обработанных источников.",
        "out/results: сформированные Excel-отчеты.",
        "monthly_baselines: сохраненные месячные снимки НП для расчета расхода в следующих месяцах.",
        "docs: пользовательские инструкции и история версий проекта.",
    ])
    doc.core_properties.title = f"Описание работы daily script {VERSION}"
    return doc


def user_guide() -> Document:
    doc = Document()
    setup(doc, compact=True)
    title(doc, "Инструкция по работе с проектом", f"Версия {VERSION}")
    add_text(doc, "Краткая инструкция для ежедневного запуска основной справки RAW04.", True)

    doc.add_heading("1 Первый запуск", level=1)
    add_numbered(doc, [
        "Поместите report_builder 2.7.0.exe в отдельную рабочую папку.",
        "Запустите программу один раз. Рядом создаются папки in, out и другие рабочие папки.",
        "Закройте программу и положите актуальную большую выгрузку в папку in.",
    ])

    doc.add_heading("2 Ежедневная справка RAW04", level=1)
    add_numbered(doc, [
        "Оставьте в папке in большую выгрузку. Ее имя не важно: программа определяет файл по колонкам.",
        "При наличии добавьте report_*.xlsx и файл готовности АПП ЖД Территория.xlsx. Они не обязательны.",
        "Запустите EXE и в главном меню введите 1.",
        "Дождитесь строк [main] RAW04 output created и [main] done.",
        "Откройте готовый файл в out/results.",
    ])
    add_text(doc, "Если report отсутствует, программа не прерывается. Таблица8 сообщит, что остатки и расход НП не рассчитаны. Это нормальное предупреждение, а не ошибка.", True)

    doc.add_heading("3 Остальные режимы", level=1)
    add_table(doc, ["Как открыть", "Пункт", "Режим"], [
        ["Главное меню 2", "1", "Старый метод, три привычных файла."],
        ["Главное меню 2", "2", "Новый метод, большая выгрузка и файл готовности."],
        ["Главное меню 2", "3", "Режим v2.7.0: с АПП КГД и большой выгрузкой."],
        ["Главное меню 2", "4", "Смешанный отчет."],
        ["Главное меню 2", "5", "Единый период с 11 февраля 2026 года."],
        ["Главное меню 3", "-", "Выход."],
    ], [3.6, 1.2, 10.7])

    doc.add_heading("4 Важные правила", level=1)
    add_bullets(doc, [
        "Закрывайте все Excel-файлы перед запуском программы.",
        "Не кладите несколько одинаковых входных файлов в папку in.",
        "Не удаляйте monthly_baselines: в ней лежит история для месячного расхода НП.",
        "Для переноса на другой компьютер копируйте всю рабочую папку программы, включая in, out и monthly_baselines.",
        "Проверяйте Диагностику, если посты или НП не распределились ожидаемым образом.",
    ])
    doc.add_heading("5 Где искать результат", level=1)
    add_table(doc, ["Папка", "Результат"], [
        ["out/results", "Excel-отчет ежедневной справки RAW04."],
        ["in/processed", "Архив входных файлов после успешного завершения."],
        ["monthly_baselines", "Снимки НП для расчетов в новых месяцах."],
    ], [4.3, 11.2])
    doc.core_properties.title = f"Инструкция по работе с проектом {VERSION}"
    return doc


def main() -> None:
    targets = [
        (ROOT / "Описание работы daily_script.docx", technical()),
        (ROOT / f"Инструкция по работе с проектом v{VERSION}.docx", user_guide()),
    ]
    for path, document in targets:
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
