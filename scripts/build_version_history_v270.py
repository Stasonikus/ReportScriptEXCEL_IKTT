from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "История версий проекта.docx"
FONT = "Times New Roman"
BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
BORDER = "D9D9D9"


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    tc_pr.append(item)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    items = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), BORDER)
        items.append(edge)
    tc_pr.append(items)


def margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    items = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), "110")
        edge.set(qn("w:type"), "dxa")
        items.append(edge)
    tc_pr.append(items)


def paragraph(doc: Document, text: str, *, note: bool = False) -> None:
    p = doc.add_paragraph(style="Document note" if note else "Normal")
    p.add_run(text)


def add_table(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (2.2, 3.1, 10.1)
    headers = ("Версия", "Статус", "Изменения")
    for col, value in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.width = Cm(widths[col])
        cell.text = value
        shade(cell, BLUE)
        borders(cell)
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, 10, True)
                run.font.color.rgb = RGBColor(255, 255, 255)

    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            cell.width = Cm(widths[col])
            cell.text = value
            if index % 2:
                shade(cell, PALE_BLUE)
            borders(cell)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col < 2 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, 9.5, col == 0)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)

    for name, size in (("Heading 1", 15), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    note = doc.styles.add_style("Document note", WD_STYLE_TYPE.PARAGRAPH)
    note.font.name = FONT
    note.font.size = Pt(10.5)
    note.font.italic = True
    note._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    note.paragraph_format.space_after = Pt(6)

    title_p = doc.add_paragraph(style="Title")
    title_p.add_run("История версий проекта")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("ReportScriptEXCEL IKTT")
    set_run_font(run, 14, True)

    paragraph(
        doc,
        "Документ фиксирует известные изменения проекта и правила нумерации сборок. "
        "Актуальная версия: 2.7.0.",
        note=True,
    )

    doc.add_heading("Правило нумерации версий", level=1)
    paragraph(doc, "Версия имеет вид X.Y.Z.")
    rules = [
        ("Первая цифра X", "Номер этапа проекта. Меняется при переходе на новый этап реализации или при смене базовой логики проекта."),
        ("Вторая цифра Y", "Масштабное изменение отчета: новые таблицы, новый режим, новая структура входных или выходных данных."),
        ("Третья цифра Z", "Микрообновление в рамках текущей версии отчета: исправление фильтра, шаблона, пункта пропуска, формата или текста меню."),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (label, text) in enumerate(rules):
        cells = table.add_row().cells
        for col, value in enumerate((label, text)):
            cell = cells[col]
            cell.width = Cm(4.3 if col == 0 else 11.1)
            cell.text = value
            shade(cell, PALE_BLUE if index % 2 else "FFFFFF")
            borders(cell)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, 10.5, col == 0)

    doc.add_heading("Журнал версий", level=1)
    paragraph(
        doc,
        "Для ранних архивных сборок, по которым не сохранился отдельный журнал изменений, указано назначение версии без детализации. "
        "Подробно зафиксированы изменения, внедренные в актуальной серии 2.6 и 2.7.",
        note=True,
    )
    add_table(doc, [
        ("2.2.1", "Архивная", "Одна из ранних распространявшихся сборок проекта. Отдельный журнал изменений не сохранен."),
        ("2.3.1", "Архивная", "Ранняя стабилизированная сборка. Отдельный журнал изменений не сохранен."),
        ("2.4.2", "Архивная", "Добавлена обработка вариантов заголовков большой выгрузки: «Водителей», «Путевых точек», «Операторов», «Промежуточные пункты»."),
        ("2.4.4-2.4.6", "Архивная", "Серия обновлений рабочего процесса, инструкций и структуры ежедневного отчета. Версия 2.4.6 сопровождалась актуальной пользовательской инструкцией."),
        ("2.5.0", "Архивная", "Промежуточная сборка перед серией 2.6. Детальный отдельный журнал не сохранен."),
        ("2.6.0-2.6.5", "Развитие", "Развитие отчетов второго этапа, разделения авто и ЖД, реестра пунктов, сохранения месячных снимков НП, Power BI и режима ежедневной справки RAW04."),
        ("2.6.6", "Исправление", "В RAW04 уточнена Таблица3: учитываются активированные и завершенные перевозки прослеживания. В Таблице6 закреплены статичные показатели пилота РК-РК."),
        ("2.6.7", "Исправление", "Таблица4 считает завершенные перевозки по направлениям из РК и в РК без фильтра по владельцу, отдельно для авто и ЖД."),
        ("2.6.8", "Функция", "Добавлен лист «Шапка» перед Таблицей1. Он содержит показатели справочного текста в фиксированном порядке и с форматированием чисел."),
        ("2.6.9", "Исправление", "RAW04 сделан запускаемым только по большой выгрузке: отсутствие report и файла готовности стало предупреждением. При отсутствии report Таблица8 показывает, что остатки и месячный расход недоступны."),
        ("2.7.0", "Текущая", "RAW04 перенесен в главное меню как основной режим. Прежний основной расчет перенесен в дополнительные режимы под названием «Режим v2.7.0: с АПП КГД и большой выгрузкой»."),
    ])

    doc.add_heading("Текущий порядок работы", level=1)
    paragraph(
        doc,
        "В версии 2.7.0 пункт 1 главного меню запускает ежедневную справку RAW04. Для нее обязательна только большая выгрузка. "
        "Файл report и файл готовности АПП/ЖД используются как дополнительные источники: программа сообщает об их отсутствии, но не прекращает расчет.",
    )
    paragraph(
        doc,
        "Пункт 2 открывает дополнительные режимы. В нем пункт 3 запускает расчет с АПП КГД и большой выгрузкой. "
        "Остальные режимы и их расчеты не переименовывают и не заменяют основной RAW04.",
    )

    props = doc.core_properties
    props.title = "История версий проекта"
    props.subject = "Журнал изменений ReportScriptEXCEL IKTT"
    props.author = "ReportScriptEXCEL IKTT"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
