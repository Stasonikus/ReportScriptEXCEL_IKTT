from __future__ import annotations

import datetime as dt
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
VERSION = "2.4.6"
UPDATED = "08.10.2026"
FONT = "Times New Roman"
BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_BLUE = "D9EAF7"
BORDER = "D9D9D9"


def _set_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_borders(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _configure_document(doc: Document, *, compact: bool = False) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8 if compact else 2.0)
    section.bottom_margin = Cm(1.8 if compact else 2.0)
    section.left_margin = Cm(2.5 if compact else 2.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12 if compact else 13)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.0)
    normal.paragraph_format.line_spacing = 1.08 if compact else 1.12
    normal.paragraph_format.space_after = Pt(4)

    title = styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(19 if compact else 21)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 15 if compact else 16, 11, 5),
        ("Heading 2", 13 if compact else 14, 8, 4),
        ("Heading 3", 12 if compact else 13, 6, 3),
    ):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Document note" not in styles:
        note = styles.add_style("Document note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = styles["Document note"]
    note.font.name = FONT
    note.font.size = Pt(10.5 if compact else 11.5)
    note.font.italic = True
    note.font.color.rgb = RGBColor(70, 70, 70)
    note._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.space_after = Pt(4)

    if "Table text" not in styles:
        table_text = styles.add_style("Table text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_text = styles["Table text"]
    table_text.font.name = FONT
    table_text.font.size = Pt(9.3 if compact else 10)
    table_text._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    table_text.paragraph_format.first_line_indent = Cm(0)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.style = table_text
    _page_number(footer)


def _title_block(doc: Document, title: str, subtitle: str, *, cover: bool = False) -> None:
    if cover:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Cm(3.2)
    paragraph = doc.add_paragraph(style="Title")
    paragraph_ppr = paragraph._p.get_or_add_pPr()
    paragraph_border = paragraph_ppr.find(qn("w:pBdr"))
    if paragraph_border is not None:
        paragraph_ppr.remove(paragraph_border)
    paragraph.add_run(title)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run(subtitle)
    _set_font(run, 15, True)


def _meta_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, values in zip(table.rows, rows):
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.text = value
            cell.width = Cm(5.0 if index == 0 else 10.0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            _set_cell_borders(cell)
            if index == 0:
                _set_cell_shading(cell, LIGHT_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles["Table text"]
                if index == 0:
                    paragraph.runs[0].bold = True


def _paragraph(doc: Document, text: str, *, note: bool = False, no_indent: bool = False) -> None:
    paragraph = doc.add_paragraph(style="Document note" if note else "Normal")
    if no_indent:
        paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.add_run(text)


def _numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, 1):
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.hanging_indent = Cm(0.55)
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(item)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.hanging_indent = Cm(0.45)
        paragraph.add_run("• ")
        paragraph.add_run(item)


def _table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    *,
    font_size: float | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    header = table.rows[0]
    _repeat_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value
        _set_cell_shading(cell, BLUE)
        _set_cell_margins(cell)
        _set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Cm(widths[index])
        for paragraph in cell.paragraphs:
            paragraph.style = doc.styles["Table text"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                if font_size:
                    run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = str(value)
            _set_cell_margins(cell)
            _set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[index])
            if row_index % 2:
                _set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles["Table text"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if font_size:
                    for run in paragraph.runs:
                        run.font.size = Pt(font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def _page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _set_properties(doc: Document, title: str, comment: str) -> None:
    properties = doc.core_properties
    properties.title = title
    properties.subject = "Автоматическое формирование отчетности по навигационным пломбам"
    properties.author = "ReportScriptEXCEL_IKTT"
    properties.comments = comment
    properties.modified = dt.datetime(2026, 10, 8, 0, 0, 0)


def build_technical_description() -> Document:
    doc = Document()
    _configure_document(doc)
    _title_block(
        doc,
        "Описание работы daily script",
        "Фактическая логика формирования отчетов по навигационным пломбам",
        cover=True,
    )
    _meta_table(
        doc,
        [
            ("Версия проекта", VERSION),
            ("Дата актуализации", UPDATED),
            ("Начало второго этапа", "01.09.2026 00:00"),
            ("Результаты запуска", "Служебный Excel-отчет и файл для Power BI"),
        ],
    )
    _paragraph(
        doc,
        "Документ описывает текущую реализацию проекта: меню программы, входные источники, правила фильтрации, распределение авто и ЖД, шесть аналитических таблиц, месячный расход НП и новый формат данных Power BI.",
        note=True,
    )
    _paragraph(doc, "Республика Казахстан, 2026", note=True)
    _page_break(doc)

    doc.add_heading("1 Общие сведения", level=1)
    _paragraph(
        doc,
        f"Программа версии {VERSION} автоматически формирует ежедневную отчетность по перевозкам и навигационным пломбам. Она принимает несколько вариантов исходных выгрузок, нормализует страны, статусы, пункты и транспорт, удаляет повторные версии перевозок, применяет правила второго этапа и создает два согласованных Excel-файла.",
    )
    _paragraph(
        doc,
        "Второй этап на территории Республики Казахстан начинается 1 сентября 2026 года в 00:00. В старом накопительном отчете результаты второго этапа получают вычитанием фиксированной базы первого этапа по состоянию на 31 августа 2026 года. В большой выгрузке записи второго этапа отбираются напрямую по полю «Дата создания».",
    )

    doc.add_heading("2 Рабочие папки и встроенные данные", level=1)
    _table(
        doc,
        ["Папка или файл", "Назначение"],
        [
            ["in", "Входные файлы текущего запуска."],
            ["in/processed", "Архив источников, использованных в успешно завершенном запуске."],
            ["out/results", "Форматированные отчеты с листами Таблица1-Таблица6."],
            ["out/power bi data", "Отдельные наборы данных для Power BI."],
            ["monthly_baselines", "Изменяемая история НП. Папка находится рядом с EXE и папками in и out; содержит снимки по датам report."],
            ["Встроенные данные EXE", "База первого этапа по 31.08.2026, стартовая месячная база на 30.09.2026, реестр пунктов, координаты и справочники."],
        ],
        [4.0, 11.8],
    )
    _paragraph(
        doc,
        "Папки создаются автоматически в корне программы, рядом с EXE. Входные файлы перемещаются в in/processed только после успешной записи служебного отчета и файла Power BI. Одинаковые архивные имена не перезаписываются: программа добавляет временную метку и при необходимости счетчик.",
    )
    _paragraph(
        doc,
        "При переносе программы на другой компьютер необходимо копировать весь корень: EXE, in, out и monthly_baselines. Замена только EXE не удаляет историю, если новый файл помещен в прежнюю рабочую папку. Удаление monthly_baselines удаляет накопленную пользовательскую историю, но не затрагивает базы, встроенные в EXE.",
    )

    doc.add_heading("3 Входные файлы", level=1)
    doc.add_heading("3 1 Обычный отчет report", level=2)
    _paragraph(
        doc,
        "Файл должен соответствовать маске report_*.xlsx и содержать лист «Отчёт». Он обязателен для Основного режима, старого метода и смешанного метода. В Основном режиме из него формируется Таблица5, а также месячный снимок НП. Файл, распознанный как большая выгрузка, не используется одновременно как обычный report.",
    )
    doc.add_heading("3 2 Большая выгрузка", level=2)
    _paragraph(
        doc,
        "Имя файла не фиксировано. Источник определяется по набору заголовков в первых 50 строках. Поддерживаются XLSX и HTML-таблица с расширением XLS. Бинарный XLS без структуры большой выгрузки рассматривается как возможная ЖД-выгрузка.",
    )
    _table(
        doc,
        ["Группа", "Обязательные поля большой выгрузки"],
        [
            ["Идентификация", "Дата создания; Номер перевозки; Статус; Дата обновления"],
            ["Классификация", "Тип перевозки; Тип транспорта; Владелец; Компания; БИН"],
            ["ТС и пломбы", "Номера ТС; Номера НП; Коды ТН ВЭД; Номера документов"],
            ["Маршрут", "Начальный пункт; Страна начала; Конечный пункт; Страна завершения; Промежуточные точки; Промежуточные страны"],
            ["Счетчики", "Кол. путевых точек; Кол. водителей; Кол. операторов"],
        ],
        [4.0, 11.8],
    )
    _paragraph(
        doc,
        "Также принимаются новые подписи «Путевых точек», «Водителей», «Операторов» и «Промежуточные пункты». Они автоматически преобразуются в канонические названия. Вариант «Железнодорожный траспорт» с системной опечаткой распознается как ЖД во всех фильтрах большой выгрузки.",
    )
    doc.add_heading("3 3 Файл готовности", level=2)
    _paragraph(
        doc,
        "Поддерживается XLSX, имя которого начинается со слова «Готовность» и содержит признаки АПП и Территория. Программа ищет лист со словами «распределение» и «НП», а при его отсутствии берет первый лист. Из нескольких датированных колонок выбирается наиболее поздняя фактическая дата не позднее текущего дня; при одинаковой дате предпочтение получает колонка с большим количеством заполненных пунктов.",
    )
    doc.add_heading("3 4 Железнодорожная выгрузка", level=2)
    _paragraph(
        doc,
        "Поддерживаются «ЖД Выгрузка*.xlsx», бинарный XLS и HTML-XLS. Источник используется в старом и смешанном методах. Для обработки требуются поля «Дата», «Владелец записи», «Уникальный номер перевозки», «Статус», company_id, start waypoint и end waypoint; для экспорта и взаимной торговли также используется поле «Тип».",
    )
    _paragraph(
        doc,
        "В папке in должен находиться не более чем один файл каждого требуемого типа. Временные файлы Excel с префиксом ~$ игнорируются.",
        note=True,
    )

    doc.add_heading("4 Меню и режимы обработки", level=1)
    _table(
        doc,
        ["Пункт", "Главное меню", "Действие"],
        [
            ["1", "Основной режим", "Запускает внутренний режим 3: большая выгрузка с разделением авто/ЖД, report и файл готовности."],
            ["2", "Остальные режимы", "Открывает меню старого, нового и смешанного методов."],
            ["3", "Выход", "Завершает программу без обработки."],
        ],
        [1.3, 4.0, 10.5],
    )
    _table(
        doc,
        ["Пункт", "Режим", "Входные файлы", "Фактические источники таблиц"],
        [
            ["0", "Назад", "Нет", "Возврат в главное меню."],
            ["1", "Старый метод", "report обязателен; готовность и ЖД-выгрузка рекомендуются", "Т1-Т2, Т4-Т6 из старых источников; Т3 - пустой шаблон. Для накопительных показателей применяется база 31.08.2026."],
            ["2", "Новый метод", "Большая выгрузка обязательна; готовность рекомендуется", "Т1-Т6 рассчитываются по большой выгрузке; файл готовности дает остатки НП для Т5."],
            ["4", "Смешанный метод", "Большая выгрузка, report, готовность и ЖД-выгрузка", "Т1-Т3 из большой выгрузки с ЖД-уточнением; Т4-Т6 из старого метода."],
        ],
        [1.0, 3.0, 5.2, 7.0],
        font_size=8.8,
    )
    _paragraph(
        doc,
        "Основной режим не читает отдельную ЖД-выгрузку. Авто и ЖД в нем определяются по большой выгрузке, форматам номеров ТС и реестру пунктов. Для Основного режима обязательны большая выгрузка, report и файл готовности.",
    )

    doc.add_heading("5 Общие правила большой выгрузки", level=1)
    _numbered(
        doc,
        [
            "Проверяется полный набор обязательных заголовков. Имя файла не влияет на распознавание.",
            "Поле «Дата создания» преобразуется в дату и время. В расчеты второго этапа входят записи с датой не ранее 01.09.2026 00:00.",
            "Строки группируются по номеру перевозки. Остается запись с самой поздней датой обновления, а при равенстве - последняя строка файла.",
            "Номера НП из всех дублей одной перевозки объединяются в множество уникальных номеров.",
            "Страны приводятся к KZ, RU, BY, KG и ARM. Пустая или некорректная дата создания не заменяется текущей датой и не попадает во второй этап.",
            "Статусы, виды перевозок, названия пунктов, владельцы и номера ТС нормализуются до подсчета.",
        ],
    )

    doc.add_heading("6 Определение транспорта и пункта", level=1)
    doc.add_heading("6 1 Авто и ЖД", level=2)
    _paragraph(
        doc,
        "В режимах с большой выгрузкой транспорт определяется по приоритету: формат «Номера ТС» из vehicle_plate_formats.json; подтверждение номера перевозки ЖД-выгрузкой в смешанном режиме; значение «Тип транспорта»; при отсутствии классификации - автомобильная группа для совместимости. Если в поле несколько номеров, автомобильный шаблон проверяется раньше железнодорожного.",
    )
    _paragraph(
        doc,
        "В Таблице1 окончательное размещение контролирует реестр пунктов. Пункт, зарегистрированный только как автомобильный или только как железнодорожный, всегда попадает в соответствующий раздел.",
    )
    doc.add_heading("6 2 Пункты Таблицы1", level=2)
    _table(
        doc,
        ["Автомобильные пункты", "Железнодорожные и стыковочные пункты"],
        [[
            "Алаколь; Казыгурт; Нур жолы; Калжат; Бахты; Тажен; Атамекен; Б. Конысбаева; Майкапчагай; Темир-баба; Капланбек",
            "Морпорт Актау; Курык порт; Болашак; Оазис; Сарыагаш; Алтынколь; Темиржол",
        ]],
        [8.0, 8.0],
    )
    _paragraph(
        doc,
        "Капланбек расположен в автомобильном разделе после Темир-баба. Темиржол расположен в ЖД-разделе. Морпорт Актау и Курык порт считаются железнодорожными или стыковочными пунктами. Для сопоставления используются канонические названия, aliases и варианты для конкретного источника из border_crossing_points.json. Нераспознанные начальные пункты учитываются в диагностике unresolved_t1, но не распределяются по строкам Таблицы1.",
    )

    doc.add_heading("7 Формирование шести таблиц", level=1)
    doc.add_heading("7 1 Таблица1 транзит с наложением НП", level=2)
    _paragraph(
        doc,
        "Для прямого расчета из большой выгрузки одновременно требуются: дата создания от 01.09.2026; статус «Активированная перевозка» или «Завершенная перевозка»; тип перевозки «транзит»; страна начала Казахстан; страна завершения Россия, Беларусь, Кыргызстан или Армения; распознанный начальный пункт. Страна завершения определяет столбец РФ, РБ, КР или РА. В Основном и смешанном режимах таблица разделена на авто и ЖД. Пустые значения отображаются тире.",
    )
    _paragraph(
        doc,
        "«Количество оформленных ЖД» в Основном режиме считается напрямую из большой выгрузки: ЖД, статус активирован или завершен, начало в РК, допустимый владелец. В старом и смешанном методах ЖД-выгрузка дополнительно применяет дату от 01.09.2026, допустимых владельцев, статусы activated/finished и исключает company_id = 1.",
    )
    doc.add_heading("7 2 Таблица2 экспорт и взаимная торговля", level=2)
    _paragraph(
        doc,
        "Учитываются уникальные перевозки с датой создания от 01.09.2026, статусом «Активированная перевозка» или «Завершенная перевозка», началом в РК и завершением в РФ, РБ, КР или РА. Тип перевозки должен быть экспортом либо взаимной торговлей. В Основном и смешанном режимах выводятся отдельные блоки авто и ЖД.",
    )
    doc.add_heading("7 3 Таблица3 услуги прослеживания", level=2)
    _paragraph(
        doc,
        "Таблица показывает перевозки, которые проходят через Казахстан, но не начинаются и не заканчиваются в Казахстане. Фильтры: дата создания от 01.09.2026; статус активирован или завершен; тип перевозки транзит; одно из четырех направлений РБ-КР, РФ-КР, КР-РФ или КР-РБ. Значения разделяются на авто и ЖД. В старом методе, где большой выгрузки нет, создается пустой шаблон Таблицы3.",
    )
    doc.add_heading("7 4 Таблица4 завершенные перевозки", level=2)
    _paragraph(
        doc,
        "После добавления новой Таблицы3 прежняя таблица завершенных перевозок сдвинута на четвертую позицию. Она содержит четыре блока: авто из РК, авто в РК, ЖД из РК и ЖД в РК. Для выезда из РК в прямом расчете используется статус «Завершенная перевозка», целевая страна и допустимый владелец. Для въезда в РК принимаются статусы «Завершенная перевозка» и «Деактивированная перевозка», начало в одной из четырех стран и завершение в Казахстане.",
    )
    doc.add_heading("7 5 Таблица5 остатки и использование НП", level=2)
    _paragraph(
        doc,
        "Левая таблица показывает текущий остаток НП из файла готовности. Итог дополнительно учитывает статичное значение 391 для пилота Карасу. Эти остатки являются снимком и не уменьшаются на базу первого этапа. Средняя таблица показывает накопительное использование НП второго этапа. В Основном режиме она формируется только из report и уменьшается на базу 31.08.2026; данные большой выгрузки сюда не подмешиваются.",
    )
    _paragraph(
        doc,
        "Правая таблица «Ежедневный расход пломб» рассчитывается отдельно. Использование в текущем месяце равно текущему накопительному значению минус снимок на последний день предыдущего месяца. После успешного формирования отчета программа сохраняет снимок текущей даты в корневой папке monthly_baselines. При переходе месяца сначала ищется точный снимок последнего календарного дня, затем самый поздний доступный снимок предыдущего месяца. Во втором случае программа выводит предупреждение.",
    )
    _paragraph(
        doc,
        "Для октября 2026 года база на 30.09.2026 с итогом 55 564 НП встроена в EXE, поэтому первый запуск на новом компьютере не требует внешнего JSON. Если пользовательский снимок поврежден, программа пропускает его и использует исправную встроенную копию, когда она существует. При полном отсутствии подходящей базы месячный расход отображается тире. Средний расход равен месячному использованию, деленному на номер текущего дня месяца с округлением до целого.",
    )
    doc.add_heading("7 6 Таблица6 активные перевозки в сторону РК", level=2)
    _paragraph(
        doc,
        "После сдвига прежняя Таблица5 стала Таблицей6. Для большой выгрузки учитываются уникальные перевозки с датой создания от 01.09.2026, статусом «Активированная перевозка», началом в РФ, РБ, КР или РА и завершением в Казахстане. В старом методе применяется прежний блок report без дополнительного вычитания.",
    )

    doc.add_heading("8 База второго этапа", level=1)
    _paragraph(
        doc,
        "Встроенная база stage2_baseline_20260831.json закрывает период первого этапа с 11 февраля по 31 августа 2026 года и имеет контрольную границу 01.09.2026 00:00. Для накопительных таблиц старого report используется формула Z = X - Y, где X - текущее накопительное значение, Y - зафиксированное значение на 31.08.2026, Z - результат второго этапа. Пользователю не нужно передавать этот JSON отдельно: при сборке он включается в EXE.",
    )
    _table(
        doc,
        ["Показатель продления первого этапа", "Зафиксированное значение"],
        [
            ["Оформлено перевозок", "6 961"],
            ["Транзит", "6 755: авто 6 608, ЖД 147"],
            ["Экспорт", "70: авто 37, ЖД 33"],
            ["Взаимная торговля", "136: авто 133, ЖД 3"],
            ["Мониторинг через РК", "1 940: авто 1 915, ЖД 25"],
            ["Завершено на территории РК", "9 936: авто 9 361, ЖД 575"],
        ],
        [8.0, 8.0],
    )
    _paragraph(
        doc,
        "Новая Таблица3 и показатели, рассчитанные напрямую по большой выгрузке с датой от 01.09.2026, повторно на базу первого этапа не уменьшаются.",
    )

    doc.add_heading("9 Выходной файл Power BI", level=1)
    _paragraph(
        doc,
        "Одновременно со служебным отчетом программа формирует отдельный Excel-файл в out/power bi data. Названия листов и столбцов являются частью фиксированного формата и не должны изменяться вручную.",
    )
    _table(
        doc,
        ["Лист", "Содержание"],
        [
            ["01_Перевозки", "Плоские ненулевые факты Таблиц1-Таблица3: дата среза, этап, показатель, направление, транспорт, пункт, страна, количество и источник."],
            ["02_Пломбы", "Остаток, накопительное использование, использование текущего месяца и примечание по пунктам."],
            ["03_Страны", "Справочник кодов RF, RB, KG и AM с русскими названиями и сокращениями."],
            ["04_Места (3)", "Фиксированный справочник мест с широтой, долготой, типом пункта, комментарием и полем Пункт_перевозки."],
            ["Количество перевозок", "Контрольная сумма строк фактов перевозок. Это сумма выгруженных показателей, а не distinct count всех исходных перевозок."],
        ],
        [4.3, 11.5],
    )
    _paragraph(
        doc,
        "Служебный отчет и Power BI-файл создаются в одном запуске и получают одинаковую временную метку. Основной отчет не изменяется ради Power BI: файл для Power BI является отдельным результатом.",
    )

    doc.add_heading("10 Имена результатов и архивирование", level=1)
    _table(
        doc,
        ["Режим", "Суффикс имени результата"],
        [
            ["Старый метод", "без дополнительного суффикса"],
            ["Новый метод", "new method"],
            ["Основной режим", "combined method"],
            ["Смешанный метод", "hybrid method"],
        ],
        [6.0, 10.0],
    )
    _paragraph(
        doc,
        "Шаблон имени: stage 2 report [суффикс]_ГГГГММДД_ЧЧММ.xlsx. После успешной записи обоих результатов использованные источники перемещаются в in/processed. Для режимов с report сохраняется новый JSON-снимок использования НП в корневой папке monthly_baselines. Имя снимка имеет вид np_snapshot_ГГГГММДД.json; повторная обработка той же даты безопасно обновляет этот файл.",
    )

    doc.add_heading("11 Диагностика", level=1)
    _table(
        doc,
        ["Сообщение или показатель", "Значение и действие"],
        [
            ["Большая выгрузка не найдена", "Нет единственного XLSX/HTML-XLS с полным набором заголовков."],
            ["Найдено несколько файлов", "Оставить в in только один источник соответствующего типа."],
            ["invalid_creation_dates", "Число строк с пустой или нераспознанной датой создания; они исключены из второго этапа."],
            ["unresolved_t1", "Число перевозок, начальный пункт которых не сопоставлен с реестром."],
            ["unresolved_np", "Количество НП, отнесенных к строке «Навешивания на территории РК»."],
            ["duplicate_rows", "Количество повторных строк перевозок, сведенных к последней версии."],
            ["База предыдущего месяца не найдена", "Проверить monthly_baselines. Для октября 2026 программа должна использовать встроенную базу 30.09.2026; для следующих месяцев нужен хотя бы один сохраненный снимок предыдущего месяца."],
            ["Не удалось записать результат", "Закрыть выходной Excel, проверить права записи и свободное место."],
        ],
        [5.2, 10.6],
    )

    doc.add_heading("12 Контрольные правила версии 2 4 6", level=1)
    _bullets(
        doc,
        [
            "Граница второго этапа - 01.09.2026 00:00, фильтр большой выгрузки - только по «Дате создания».",
            "Таблица3 - проследование через РК; прежние Таблицы3, 4 и 5 сдвинуты на позиции 4, 5 и 6.",
            "Основной режим запускается пунктом 1 главного меню и не требует отдельной ЖД-выгрузки.",
            "Ошибочное значение «Железнодорожный траспорт» распознается как ЖД.",
            "Нулевые ячейки в форматированных таблицах отображаются тире там, где это предусмотрено шаблоном.",
            "Капланбек - авто; Темиржол, Морпорт Актау и Курык порт - ЖД или стыковочные пункты.",
            "Накопительное использование НП второго этапа и расход НП текущего месяца рассчитываются по разным базам.",
            "База первого этапа по 31.08.2026 и месячная база на 30.09.2026 встроены в EXE.",
            "Новые месячные снимки сохраняются рядом с EXE в monthly_baselines и переносятся вместе со всей папкой программы.",
            "Формат Power BI содержит пять листов с фиксированными названиями и отдельный координатный справочник мест.",
        ],
    )
    _paragraph(doc, "Конец документа.", note=True)
    _set_properties(
        doc,
        f"Описание работы daily script {VERSION}",
        "Обновлено по фактической логике проекта: граница 01.09.2026, шесть таблиц, новая Таблица3 и новый Power BI-файл.",
    )
    return doc


def build_full_user_guide() -> Document:
    doc = Document()
    _configure_document(doc, compact=True)
    _title_block(doc, "Инструкция по работе с проектом", f"Версия {VERSION}  Полное руководство пользователя")
    _paragraph(
        doc,
        "Инструкция предназначена для ежедневного запуска готового EXE-файла. Основной рабочий сценарий использует большую выгрузку, обычный report и файл готовности, формирует шесть листов отчета и отдельный файл Power BI.",
        note=True,
    )

    doc.add_heading("1 Что подготовить для Основного режима", level=1)
    _table(
        doc,
        ["Файл", "Требование", "Для чего используется"],
        [
            ["Большая выгрузка", "Один XLSX или HTML-XLS с колонками «Дата создания», «Номер перевозки» и остальными обязательными полями", "Таблицы1-Таблица4 и Таблица6, разделение авто/ЖД"],
            ["report_*.xlsx", "Один файл с листом «Отчёт»", "Накопительное использование НП и месячная база Таблицы5"],
            ["Готовность АПП_ЖД_Территория*.xlsx", "Один XLSX", "Текущий остаток НП в левой части Таблицы5"],
        ],
        [4.2, 6.0, 5.6],
        font_size=8.8,
    )
    _paragraph(
        doc,
        "Отдельная ЖД-выгрузка для Основного режима не нужна. Закройте все входные и ранее созданные Excel-файлы перед запуском.",
    )

    doc.add_heading("2 Первый запуск", level=1)
    _numbered(
        doc,
        [
            "Поместите актуальный report_builder.exe в отдельную рабочую папку.",
            "Запустите EXE один раз. Рядом появятся папки in, out и необходимые подпапки. Папка monthly_baselines создается при обработке report.",
            "Закройте программу и скопируйте входные файлы в папку in.",
            "Не помещайте исходники в in/processed. Эта папка заполняется программой автоматически.",
        ],
    )

    doc.add_heading("3 Ежедневный запуск", level=1)
    _numbered(
        doc,
        [
            "Удалите из in остатки неудачного предыдущего запуска, но не удаляйте папку processed.",
            "Скопируйте в in по одному актуальному файлу каждого требуемого типа.",
            f"Запустите report_builder.exe. В консоли должна появиться версия {VERSION}.",
            "В главном меню нажмите 1 - «Основной режим».",
            "Дождитесь сообщения [main] done. Не закрывайте окно во время чтения и записи файлов.",
            "Откройте отчет в out/results и файл Power BI в out/power bi data.",
            "Проверьте, что входные файлы перемещены в in/processed, а для report создан месячный снимок в monthly_baselines.",
        ],
    )

    doc.add_heading("4 Что делает Основной режим", level=1)
    _table(
        doc,
        ["Лист", "Содержание", "Основной источник"],
        [
            ["Таблица1", "Транзит из РК по пунктам, странам и авто/ЖД; количество оформленных ЖД", "Большая выгрузка"],
            ["Таблица2", "Экспорт и взаимная торговля из РК по странам и авто/ЖД", "Большая выгрузка"],
            ["Таблица3", "Проследование через РК: РБ-КР, РФ-КР, КР-РФ, КР-РБ", "Большая выгрузка"],
            ["Таблица4", "Завершенные авто и ЖД из РК и в РК", "Большая выгрузка"],
            ["Таблица5", "Остаток НП, использование второго этапа, расход текущего месяца и средний расход", "report и готовность"],
            ["Таблица6", "Активные перевозки в сторону РК", "Большая выгрузка"],
        ],
        [2.5, 8.5, 4.8],
        font_size=8.8,
    )

    doc.add_heading("5 Главное правило второго этапа", level=1)
    _paragraph(
        doc,
        "Для большой выгрузки учитываются перевозки, созданные 1 сентября 2026 года в 00:00 или позднее. Проверяется именно «Дата создания», а не «Дата обновления». Повторный номер перевозки считается один раз по самой свежей версии. Статичные накопительные показатели report уменьшаются на встроенную базу по состоянию на 31 августа 2026 года.",
    )
    _paragraph(
        doc,
        "Текущий месячный расход НП сбрасывается каждый календарный месяц: из текущего накопительного значения вычитается снимок последнего дня предыдущего месяца. Для октября 2026 используется встроенная база 30.09.2026 с итогом 55 564 НП. Для следующих месяцев программа использует снимки, автоматически сохраненные в monthly_baselines. Поэтому месячный расход не обязан совпадать с общим использованием НП второго этапа.",
    )

    doc.add_heading("6 Хранение и перенос истории", level=1)
    _paragraph(
        doc,
        "Папка monthly_baselines находится в той же рабочей папке, что EXE, in и out. После каждого успешного запуска с report программа сохраняет накопительные значения НП по каждому посту и общий итог в файл np_snapshot_ГГГГММДД.json. При наступлении нового месяца выбирается точный снимок последнего дня предыдущего месяца, а при его отсутствии - самый поздний доступный снимок того месяца с предупреждением.",
    )
    _numbered(
        doc,
        [
            "При обновлении замените только EXE и оставьте monthly_baselines на месте.",
            "При переносе на другой компьютер скопируйте всю рабочую папку программы, включая monthly_baselines.",
            "Не редактируйте и не удаляйте JSON-файлы вручную. Восстановить встроенную базу 30.09.2026 можно повторной установкой EXE, но последующие месяцы хранятся только в рабочей папке.",
        ],
    )

    _page_break(doc)
    doc.add_heading("7 Другие режимы", level=1)
    _paragraph(doc, "Нажмите 2 в главном меню. В дополнительном меню доступны:")
    _table(
        doc,
        ["Пункт", "Режим", "Необходимые файлы"],
        [
            ["0", "Назад", "Возврат в главное меню"],
            ["1", "Старый метод", "report; готовность и ЖД-выгрузка для полного результата"],
            ["2", "Новый метод", "Большая выгрузка; готовность для остатка НП"],
            ["4", "Смешанный метод", "Большая выгрузка, report, готовность и ЖД-выгрузка"],
        ],
        [1.4, 4.8, 9.6],
    )
    _paragraph(
        doc,
        "Пункт 3 главного меню завершает программу. Пункт 0 дополнительного меню только возвращает назад и не запускает расчет.",
    )

    doc.add_heading("8 Проверка результата", level=1)
    _bullets(
        doc,
        [
            "В out/results должен быть один новый файл stage 2 report...xlsx с шестью листами.",
            "В out/power bi data должен быть файл с той же временной меткой и листами 01_Перевозки, 02_Пломбы, 03_Страны, 04_Места (3), Количество перевозок.",
            "Пустые значения в форматированных таблицах должны быть обозначены тире, а не нулями, где это предусмотрено шаблоном.",
            "Капланбек должен находиться в автомобильной части Таблицы1 после Темир-баба; Темиржол, Морпорт Актау и Курык порт - в ЖД-части.",
            "В консоли проверьте invalid_creation_dates, unresolved_t1, duplicate_rows и итоговые значения таблиц.",
            "В корне программы должен появиться monthly_baselines с JSON-снимком даты обработанного report.",
        ],
    )

    doc.add_heading("9 Если программа сообщает об ошибке", level=1)
    _table(
        doc,
        ["Ситуация", "Что сделать"],
        [
            ["Файл источника не найден", "Проверить папку in, имя report_*.xlsx и выбранный режим."],
            ["Большая выгрузка не найдена", "Убедиться, что это XLSX или HTML-XLS с полным набором колонок; имя файла неважно."],
            ["Найдено несколько файлов", "Оставить по одному файлу каждого типа."],
            ["Файл готовности не найден", "Для Основного и смешанного режимов добавить единственный подходящий XLSX."],
            ["Месячная база не найдена", "Проверить, что рядом с EXE перенесена папка monthly_baselines. Для октября 2026 отдельный файл не нужен: база 30.09.2026 встроена."],
            ["Не удалось записать результат", "Закрыть открытые выходные файлы и проверить свободное место."],
        ],
        [5.2, 10.6],
    )
    _paragraph(
        doc,
        "Важно: вариант «Железнодорожный траспорт» в исходной выгрузке поддерживается и не требует ручного исправления.",
        note=True,
    )
    _set_properties(
        doc,
        f"Инструкция по работе с проектом {VERSION}",
        "Полное руководство пользователя для актуальной структуры проекта.",
    )
    return doc


def build_quick_guide() -> Document:
    doc = Document()
    _configure_document(doc, compact=True)
    _title_block(doc, "Краткая инструкция по работе с проектом", f"Версия {VERSION}  Основной режим")
    _paragraph(
        doc,
        "Основной режим формирует шесть листов служебного отчета и отдельный файл для Power BI. Расчет перевозок начинается с 01.09.2026 00:00 по полю «Дата создания».",
        note=True,
    )

    doc.add_heading("1 Подготовьте три файла", level=1)
    _table(
        doc,
        ["Файл", "Назначение"],
        [
            ["Большая выгрузка XLSX или HTML-XLS", "Таблицы1-Таблица4 и Таблица6, авто/ЖД"],
            ["report_*.xlsx", "Использование НП и текущий месяц в Таблице5"],
            ["Готовность АПП_ЖД_Территория*.xlsx", "Остаток НП в Таблице5"],
        ],
        [7.0, 8.8],
    )
    _paragraph(doc, "В папке in должен быть только один файл каждого типа. Все Excel-файлы перед запуском закройте.")

    doc.add_heading("2 Запустите программу", level=1)
    _numbered(
        doc,
        [
            "Скопируйте три файла в папку in.",
            "Запустите актуальный report_builder.exe.",
            "В главном меню нажмите 1 - «Основной режим».",
            "Дождитесь [main] done.",
        ],
    )

    doc.add_heading("3 Заберите результаты", level=1)
    _table(
        doc,
        ["Папка", "Результат"],
        [
            ["out/results", "Основной отчет с листами Таблица1-Таблица6"],
            ["out/power bi data", "Файл Power BI с пятью фиксированными листами"],
            ["in/processed", "Архив использованных входных файлов"],
            ["monthly_baselines", "Переносимая история использования НП по датам report; находится рядом с EXE"],
        ],
        [6.0, 9.8],
    )

    doc.add_heading("4 Быстрая проверка", level=1)
    _bullets(
        doc,
        [
            "Таблица3 содержит проследование через РК по четырем направлениям и авто/ЖД.",
            "Таблица5 сформирована из report и готовности; месячный расход считается отдельно от накопительного итога.",
            "Капланбек находится в авто-разделе после Темир-баба; Темиржол, Морпорт Актау и Курык порт - в ЖД-разделе.",
            "Опечатка «Железнодорожный траспорт» распознается автоматически.",
            "Power BI-файл содержит листы 01_Перевозки, 02_Пломбы, 03_Страны, 04_Места (3), Количество перевозок.",
        ],
    )

    doc.add_heading("5 Остальные команды меню", level=1)
    _table(
        doc,
        ["Команда", "Действие"],
        [
            ["Главное меню 2", "Открыть остальные режимы"],
            ["Дополнительное меню 1", "Старый метод"],
            ["Дополнительное меню 2", "Новый метод"],
            ["Дополнительное меню 4", "Смешанный метод"],
            ["Дополнительное меню 0", "Назад"],
            ["Главное меню 3", "Выход"],
        ],
        [6.0, 9.8],
    )
    _paragraph(
        doc,
        "Если программа не находит источник, сначала проверьте количество файлов в in и структуру большой выгрузки. Имя большой выгрузки не фиксировано.",
        note=True,
    )
    _set_properties(
        doc,
        f"Краткая инструкция по работе с проектом {VERSION}",
        "Краткий порядок ежедневного запуска Основного режима.",
    )
    return doc


def main() -> None:
    outputs = [
        (ROOT / "Описание работы daily_script.docx", build_technical_description()),
        (ROOT / f"Инструкция по работе с проектом v{VERSION}.docx", build_full_user_guide()),
    ]
    for path, document in outputs:
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
