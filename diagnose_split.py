from collections import Counter
import datetime as dt
from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))

from big_export import (
    COUNTED_CURRENT_STATUSES,
    COUNTRY_COLUMNS,
    REQUIRED_HEADERS,
    STAGE2_START,
    T1_POINT_ALIASES,
    _country_code,
    _find_header,
    _match_point,
    _norm,
    _parse_datetime,
    _transport,
    _transportation_type,
    extract_stage2_rail_ids,
    is_big_export_file,
    load_t1_transport_groups,
)
from config import load_config
from io_excel import read_rail_report
from utils import resource_path


input_dir = ROOT / "in"
search_dir = input_dir if any(input_dir.glob("*.xlsx")) else input_dir / "processed"
files = sorted(search_dir.glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
big_file = next((path for path in files if is_big_export_file(path)), None)
rail_file = next((path for path in files if path.name.lower().startswith("жд выгрузка")), None)
assert big_file and rail_file

rail_ids = extract_stage2_rail_ids(read_rail_report(rail_file))
cfg = load_config()
groups = load_t1_transport_groups(
    resource_path("data/border_crossing_points.json"),
    cfg.tp_zhd_list_t1,
)

book = load_workbook(big_file, read_only=True, data_only=True)
try:
    sheet = None
    header_row = None
    header_map = None
    for candidate in book.worksheets:
        found = _find_header(candidate)
        if found:
            sheet = candidate
            header_row, header_map = found
            break
    assert sheet and header_row and header_map
    indexes = {name: header_map[name] for name in REQUIRED_HEADERS}
    records = {}
    for source_row, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if not any(value not in (None, "") for value in values):
            continue
        row = {
            name: values[index] if index < len(values) else None
            for name, index in indexes.items()
        }
        created = _parse_datetime(row["Дата создания"])
        updated = _parse_datetime(row["Дата обновления"])
        key = str(row["Номер перевозки"] or "").strip() or f"__row_{source_row}"
        row["_created"] = created
        row["_updated"] = updated
        row["_source_row"] = source_row
        order = (updated or created or dt.datetime.min, source_row)
        current = records.get(key)
        if current is None:
            records[key] = row
        else:
            current_order = (
                current["_updated"] or current["_created"] or dt.datetime.min,
                current["_source_row"],
            )
            if order > current_order:
                records[key] = row
finally:
    book.close()

all_by_point = Counter()
included_by_point = Counter()
excluded_by_point = Counter()
alakol_transport_field = Counter()
alakol_membership = Counter()

for key, row in records.items():
    created = row["_created"]
    status = _norm(row["Статус"])
    movement = _transportation_type(row["Тип перевозки"])
    country_from = _country_code(row["Страна начала"])
    country_to = _country_code(row["Страна завершения"])
    if not (
        created is not None
        and created >= STAGE2_START
        and status in COUNTED_CURRENT_STATUSES
        and movement == "transit"
        and country_from == "KZ"
        and country_to in COUNTRY_COLUMNS
    ):
        continue
    point = _match_point((row["Начальный пункт"],), T1_POINT_ALIASES)
    if point is None:
        continue
    transport_group = "rail" if key in rail_ids else "auto"
    all_by_point[point] += 1
    if transport_group in groups.get(point, {"auto", "rail"}):
        included_by_point[point] += 1
    else:
        excluded_by_point[point] += 1
    if point == "Алаколь":
        alakol_transport_field[_transport(row["Тип транспорта"])] += 1
        alakol_membership[transport_group] += 1

print("BIG_FILE", big_file.name)
print("RAIL_FILE", rail_file.name)
print("ALAKOL_GROUPS", sorted(groups.get("Алаколь", set())))
print("ALAKOL_ALL", all_by_point["Алаколь"])
print("ALAKOL_INCLUDED", included_by_point["Алаколь"])
print("ALAKOL_EXCLUDED", excluded_by_point["Алаколь"])
print("ALAKOL_TRANSPORT_FIELD", dict(alakol_transport_field))
print("ALAKOL_RAIL_MEMBERSHIP", dict(alakol_membership))
print("ALL_TRANSIT_MATCHED", sum(all_by_point.values()))
print("INCLUDED_SPLIT", sum(included_by_point.values()))
print("EXCLUDED_SPLIT", sum(excluded_by_point.values()))
print("EXCLUDED_TOP", excluded_by_point.most_common(10))
