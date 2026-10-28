from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


OUTPUT_FONT_NAME = "Times New Roman"
OUTPUT_FONT_SIZE = 14
MIN_ROW_HEIGHT = 21
MAX_AUTOFIT_WIDTH = 90


def _output_font(bold: bool = False) -> Font:
    return Font(
        name=OUTPUT_FONT_NAME,
        size=OUTPUT_FONT_SIZE,
        bold=bold,
    )


def read_report_1(path: Path, sheet_name: str = "Отчёт") -> Dict[str, Any]:

    if not path.exists():
        raise FileNotFoundError(f"report_1 not found: {path}")

    wb = load_workbook(path, data_only=True)

    if sheet_name not in wb.sheetnames:
        raise KeyError(f"sheet not found: {sheet_name}")

    ws = wb[sheet_name]

    return {
        "path": str(path),
        "workbook": wb,
        "ws": ws,
        "sheet_name": sheet_name,
        "sheetnames": wb.sheetnames,
    }


def read_report_2(path: Path) -> Dict[str, Any]:

    if not path.exists():
        raise FileNotFoundError(f"report_2 not found: {path}")

    wb = load_workbook(path, data_only=True)

    target_sheet = None

    for s in wb.sheetnames:

        low = s.strip().lower()

        if "распределение" in low and "нп" in low:
            target_sheet = s
            break

    if target_sheet is None:
        target_sheet = wb.sheetnames[0]

    ws = wb[target_sheet]

    return {
        "path": str(path),
        "workbook": wb,
        "ws": ws,
        "sheet_name": target_sheet,
        "sheetnames": wb.sheetnames,
        "worksheets": wb.worksheets,
    }


def read_rail_report(path: Path) -> Dict[str, Any]:

    if not path.exists():
        raise FileNotFoundError(f"rail report not found: {path}")

    if path.suffix.lower() == ".xls":
        return _read_rail_xls_skipping_logs(path)

    wb = load_workbook(path, data_only=True)
    ws = wb.active

    return {
        "path": str(path),
        "workbook": wb,
        "ws": ws,
        "sheet_name": ws.title,
        "sheetnames": wb.sheetnames,
    }


def _read_rail_xls_skipping_logs(path: Path) -> Dict[str, Any]:

    raw = path.read_bytes()

    if raw.lstrip()[:5].lower() == b"<html":
        rows = _read_html_table_rows(raw)
        return _rail_rows_to_workbook(path, rows, sheet_name="Rail")

    import xlrd

    source_book = xlrd.open_workbook(str(path))
    source_sheet = source_book.sheet_by_index(0)

    rows = [
        [
            source_sheet.cell_value(source_row, source_col)
            for source_col in range(source_sheet.ncols)
        ]
        for source_row in range(source_sheet.nrows)
    ]

    return _rail_rows_to_workbook(path, rows, sheet_name=source_sheet.name)


class _HtmlTableParser(HTMLParser):

    def __init__(self) -> None:
        super().__init__()
        self.rows: List[List[str]] = []
        self._in_row = False
        self._in_cell = False
        self._cell_parts: List[str] = []
        self._row: List[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()

        if tag == "tr":
            self._in_row = True
            self._row = []
            return

        if tag in {"td", "th"} and self._in_row:
            self._in_cell = True
            self._cell_parts = []

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._cell_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()

        if tag in {"td", "th"} and self._in_cell:
            value = " ".join("".join(self._cell_parts).split())
            self._row.append(value)
            self._in_cell = False
            self._cell_parts = []
            return

        if tag == "tr" and self._in_row:
            if self._row:
                self.rows.append(self._row)

            self._in_row = False
            self._row = []


def _read_html_table_rows(raw: bytes) -> List[List[str]]:

    parser = _HtmlTableParser()
    parser.feed(raw.decode("utf-8", errors="replace"))

    return parser.rows


def _rail_rows_to_workbook(
    path: Path,
    rows: List[List[Any]],
    sheet_name: str,
) -> Dict[str, Any]:

    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_title(sheet_name)

    source_start_row = _find_rail_header_row_index(rows)

    if source_start_row is None:
        source_start_row = 16

    for source_row in range(source_start_row, len(rows)):
        target_row = source_row - source_start_row + 1
        row = rows[source_row]

        for source_col, value in enumerate(row):
            ws.cell(
                row=target_row,
                column=source_col + 1,
                value=value,
            )

    return {
        "path": str(path),
        "workbook": wb,
        "ws": ws,
        "sheet_name": ws.title,
        "sheetnames": [ws.title],
    }


def _find_rail_header_row_index(rows: List[List[Any]]) -> int | None:

    required = {
        "владелец записи",
        "статус",
        "company_id",
        "start waypoint",
        "end waypoint",
    }

    for idx, row in enumerate(rows):
        values = {
            str(value).strip().lower()
            for value in row
            if value is not None and str(value).strip()
        }

        if required.issubset(values):
            return idx

    return None


def write_output_test(out_path: Path, sheet_name: str = "Таблица1") -> None:

    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    ws = wb.active

    ws.title = sheet_name

    ws["A1"] = "OK"

    wb.save(out_path)


def _safe_sheet_title(name: str) -> str:

    bad = [":", "\\", "/", "?", "*", "[", "]"]

    for ch in bad:
        name = name.replace(ch, " ")

    name = name.strip()

    if len(name) > 31:
        name = name[:31]

    return name or "Sheet"


def _delete_default_sheet_if_needed(wb: Workbook) -> None:

    if len(wb.worksheets) == 1 and wb.worksheets[0].title == "Sheet":
        wb.remove(wb.worksheets[0])


def _apply_output_font_to_used_range(ws, force_bold: bool | None = None) -> None:

    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            is_bold = bool(cell.font and cell.font.bold)

            if force_bold is not None:
                is_bold = force_bold

            cell.font = _output_font(bold=is_bold)


def _is_total_value(value: Any) -> bool:

    return str(value or "").strip().lower() in {"итого", "все"}


def _bold_total_rows(ws) -> None:

    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        if any(_is_total_value(cell.value) for cell in row):
            for cell in row:
                cell.font = _output_font(bold=True)


def _set_column_bold(ws, col_idx: int) -> None:

    for row_idx in range(1, ws.max_row + 1):
        ws.cell(row=row_idx, column=col_idx).font = _output_font(bold=True)


def _fit_used_range(ws, max_width: int = MAX_AUTOFIT_WIDTH) -> None:

    for col_idx in range(1, ws.max_column + 1):
        max_line_len = 0

        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value

            if value is None:
                continue

            lines = str(value).splitlines() or [str(value)]
            max_line_len = max(max_line_len, *(len(line) for line in lines))

        if max_line_len == 0:
            continue

        letter = get_column_letter(col_idx)
        width = max(10, int(max_line_len * 1.35) + 3)
        ws.column_dimensions[letter].width = min(width, max_width)

    for row_idx in range(1, ws.max_row + 1):
        max_lines = 1

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = cell.value

            if value is None:
                continue

            text = str(value)
            explicit_lines = max(1, len(text.splitlines()))
            col_width = ws.column_dimensions[get_column_letter(col_idx)].width or 10
            wrapped_lines = max(1, int((len(text) / max(col_width, 1)) + 0.95))
            max_lines = max(max_lines, explicit_lines, wrapped_lines)

            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical=cell.alignment.vertical or "center",
                wrap_text=True,
            )

        ws.row_dimensions[row_idx].height = max(
            ws.row_dimensions[row_idx].height or 0,
            MIN_ROW_HEIGHT * max_lines,
        )


def _set_col_widths(
    ws,
    headers: List[str],
    start_col: int = 1,
    max_rows_scan: int = 200,
) -> None:

    for i, h in enumerate(headers, start=start_col):

        letter = get_column_letter(i)

        max_len = len(str(h))

        for r in range(1, min(ws.max_row, max_rows_scan) + 1):

            v = ws.cell(row=r, column=i).value

            if v is None:
                continue

            max_len = max(max_len, len(str(v)))

        ws.column_dimensions[letter].width = min(
            max(10, max_len + 2),
            55,
        )


def _write_rows_as_table(
    ws,
    headers: List[str],
    rows: List[Dict[str, Any]],
    start_row: int = 1,
    start_col: int = 1,
    dash_zero: bool = False,
) -> int:

    for j, h in enumerate(headers, start=start_col):

        c = ws.cell(row=start_row, column=j, value=h)

        c.font = _output_font(bold=True)

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    r_out = start_row + 1

    for row in rows:

        is_summary = any(_is_total_value(value) for value in row.values())

        for j, h in enumerate(headers, start=start_col):

            value = row.get(h, None)
            if dash_zero and isinstance(value, (int, float)) and value == 0:
                value = "-"

            cell = ws.cell(
                row=r_out,
                column=j,
                value=value,
            )

            if value == "-" or isinstance(value, (int, float)):
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            if is_summary:
                cell.font = _output_font(bold=True)

        r_out += 1

    return r_out


def _write_section_title(
    ws,
    title: str,
    row: int,
    start_col: int = 1,
    span_cols: int = 3,
) -> int:

    c = ws.cell(row=row, column=start_col, value=title)

    c.font = _output_font(bold=True)

    c.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    if span_cols > 1:

        ws.merge_cells(
            start_row=row,
            start_column=start_col,
            end_row=row,
            end_column=start_col + span_cols - 1,
        )

    return row + 1


def _write_table2(ws, t: Dict[str, Any]) -> None:

    export_rows = t.get("export_rows", []) or []
    trade_rows = t.get("trade_rows", []) or []

    headers = ["№", "Страна", "Количество перевозок"]

    row = 1

    row = _write_section_title(
        ws,
        "Статистические данные экспорта и взаимной торговли из РК",
        row,
        span_cols=3,
    )

    row = _write_section_title(
        ws,
        "Экспорт",
        row,
        span_cols=3,
    )

    row = _write_rows_as_table(
        ws,
        headers=headers,
        rows=export_rows,
        start_row=row,
    )

    row += 1

    row = _write_section_title(
        ws,
        "Взаимная торговля",
        row,
        span_cols=3,
    )

    row = _write_rows_as_table(
        ws,
        headers=headers,
        rows=trade_rows,
        start_row=row,
    )

    ws.freeze_panes = "A4"

    _set_col_widths(ws, headers)


def _write_table3(ws, t: Dict[str, Any]) -> None:

    auto_out_rows = t.get("auto_out_rows", []) or []
    auto_in_rows = t.get("auto_in_rows", []) or []

    rail_out_rows = t.get("rail_out_rows", []) or []
    rail_in_rows = t.get("rail_in_rows", []) or []

    headers = ["Направление", "Количество"]

    row = 1

    row = _write_section_title(
        ws,
        "Статистические данные завершенных перевозок",
        row,
        span_cols=len(headers),
    )

    row = _write_section_title(
        ws,
        "Завершённые автомобильные перевозки из Республики Казахстан",
        row,
        span_cols=len(headers),
    )

    row = _write_rows_as_table(
        ws,
        headers,
        auto_out_rows,
        row,
    )

    row += 1

    row = _write_section_title(
        ws,
        "Завершенные автомобильные перевозки в Республику Казахстан",
        row,
        span_cols=len(headers),
    )

    row = _write_rows_as_table(
        ws,
        headers,
        auto_in_rows,
        row,
    )

    row += 1

    row = _write_section_title(
        ws,
        "Завершённые железнодорожные перевозки из Республики Казахстан",
        row,
        span_cols=len(headers),
    )

    row = _write_rows_as_table(
        ws,
        headers,
        rail_out_rows,
        row,
    )

    row += 1

    row = _write_section_title(
        ws,
        "Завершенные железнодорожные перевозки в Республику Казахстан",
        row,
        span_cols=len(headers),
    )

    row = _write_rows_as_table(
        ws,
        headers,
        rail_in_rows,
        row,
    )

    ws.freeze_panes = "A4"

    _set_col_widths(ws, headers)


def _write_table4(ws, t: Dict[str, Any]) -> None:

    rows = t.get("rows", []) or []
    date_label = str(
        t.get("date_label_short") or t.get("date_label", "") or ""
    ).strip()

    headers = [
        "Пункты\nпропуска/региональ\nные склады",
        f"НП на {date_label}" if date_label else "НП",
    ]
    data_rows = [
        {
            headers[0]: row.get("Точка"),
            headers[1]: row.get("Количество НП"),
        }
        for row in rows
    ]
    total_np = sum(
        int(row.get("Количество НП") or 0)
        for row in rows
        if isinstance(row.get("Количество НП"), (int, float))
    )
    data_rows.append(
        {
            headers[0]: "Итого",
            headers[1]: total_np,
        }
    )

    row = _write_rows_as_table(
        ws,
        headers=headers,
        rows=data_rows,
        start_row=1,
        dash_zero=True,
    )

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_cells in ws.iter_rows(min_row=1, max_row=row - 1, min_col=1, max_col=2):
        for cell in row_cells:
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if cell.column == 2 else "left",
                vertical="center",
                wrap_text=True,
            )

            if cell.row == 1 or cell.column == 2:
                cell.font = _output_font(bold=True)

    used_np_rows = t.get("used_np_rows", []) or []
    used_headers = [
        "Пункты пропуска/региональные склады",
        "",
        "Кол-во\nиспольз\nованных НП",
    ]

    for col, value in zip(range(4, 7), used_headers):
        cell = ws.cell(row=1, column=col, value=value)
        cell.border = border
        cell.font = _output_font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    used_row = 2
    for item in used_np_rows:
        is_total = item.get("Пункт") in {"Навешивания на территории РК", "Итого"}

        ws.cell(row=used_row, column=4, value=item.get("Пункт"))
        ws.cell(row=used_row, column=5, value=None)
        ws.cell(row=used_row, column=6, value=item.get("Количество"))

        for col in range(4, 7):
            cell = ws.cell(row=used_row, column=col)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if col == 6 else "left",
                vertical="center",
                wrap_text=True,
            )

            if col == 6 or is_total:
                cell.font = _output_font(bold=True)

        used_row += 1

    ws.row_dimensions[1].height = 42
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 18

    ws.freeze_panes = "A2"



def _write_table5(ws, t: Dict[str, Any]) -> None:

    rows = t.get("rows", []) or []

    headers = [
        "Страна начала перевозки",
        "Количество перевозок",
    ]

    row = 1

    row = _write_section_title(
        ws,
        "Статистика по перевозкам, направленным в сторону Республики Казахстан",
        row,
        span_cols=2,
    )

    _write_rows_as_table(
        ws,
        headers=headers,
        rows=rows,
        start_row=row,
    )

    ws.freeze_panes = "A3"

    _set_col_widths(ws, headers)


def write_output(out_path: Path, tables: List[Any]) -> None:

    out_path.parent.mkdir(parents=True, exist_ok=True)

    if not tables:

        write_output_test(
            out_path,
            sheet_name="Таблица1",
        )

        print(
            f"[io_excel.write_output] "
            f"tables_stub_count=0 "
            f"(created {out_path} with Таблица1:A1='OK')"
        )

        return

    wb = Workbook()

    _delete_default_sheet_if_needed(wb)

    for t in tables:

        if not isinstance(t, dict):
            continue

        sheet_title = _safe_sheet_title(
            str(t.get("table", "Sheet"))
        )

        ws = wb.create_sheet(title=sheet_title)

        if sheet_title == "Таблица2":

            _write_table2(ws, t)
            _apply_output_font_to_used_range(ws)
            _bold_total_rows(ws)
            _fit_used_range(ws)

            continue

        if sheet_title == "Таблица3":

            _write_table3(ws, t)
            _apply_output_font_to_used_range(ws)
            _bold_total_rows(ws)
            _fit_used_range(ws)

            continue

        if sheet_title == "Таблица4":

            _write_table4(ws, t)
            _apply_output_font_to_used_range(ws, force_bold=True)
            _fit_used_range(ws)

            continue

        if sheet_title == "Таблица5":

            _write_table5(ws, t)
            _apply_output_font_to_used_range(ws)
            _bold_total_rows(ws)
            _fit_used_range(ws)

            continue

        rows = t.get("rows", [])

        if not isinstance(rows, list) or (
            rows and not isinstance(rows[0], dict)
        ):

            ws["A1"] = str(t)

            continue

        if sheet_title == "Таблица1":

            headers = [
                "ТП/ЖД",
                "РФ",
                "РБ",
                "КР",
                "РА",
                "Итого",
            ]

        else:

            headers = list(rows[0].keys()) if rows else ["data"]

        _write_rows_as_table(
            ws,
            headers,
            rows,
            dash_zero=(sheet_title == "Таблица1"),
        )

        if sheet_title == "Таблица1":
            rail_registered_total = t.get("rail_registered_total")

            title_cell = ws.cell(row=1, column=8, value="Количество оформленных ЖД")
            value_cell = ws.cell(
                row=2,
                column=8,
                value="" if rail_registered_total is None else rail_registered_total,
            )

            for cell in (title_cell, value_cell):
                cell.font = _output_font(bold=True)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            ws.column_dimensions["H"].width = 24
            _set_column_bold(ws, 6)

        ws.freeze_panes = "A2"

        _set_col_widths(ws, headers)
        _apply_output_font_to_used_range(ws)
        _bold_total_rows(ws)
        _fit_used_range(ws)

    wb.save(out_path)

    print(
        f"[io_excel.write_output] "
        f"wrote {len(tables)} table(s) to: {out_path}"
    )
