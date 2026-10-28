from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "POWER_BI_DATA_MODEL_SPEC.md"
OUTPUT = ROOT / "docs" / "Power_BI_technical_specification.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F3F7FA"
GRAY = "D9D9D9"
TEXT = RGBColor(31, 41, 55)


def set_run_font(run, name="Arial", size=10.5, bold=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_border(cell, color=GRAY, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        edge_el = tc_borders.find(qn(tag))
        if edge_el is None:
            edge_el = OxmlElement(tag)
            tc_borders.append(edge_el)
        edge_el.set(qn("w:val"), "single")
        edge_el.set(qn("w:sz"), size)
        edge_el.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_run_font(run, size=8, color=RGBColor(89, 89, 89))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.extend([r_pr, text_node])
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text, size=10.5):
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(8.5, size - 1), color=RGBColor(55, 65, 81))
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def clean_heading(text):
    text = text.replace("`", "").replace("—", " ").replace("–", " ")
    text = re.sub(r"^(\d+(?:\.\d+)*)\.\s*", lambda m: m.group(1).replace(".", " ") + " ", text)
    text = re.sub(r"[:;|]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(cols):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            value = values[c_idx] if c_idx < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if r_idx == 0:
                shade_cell(cell, NAVY)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(re.sub(r"[*`]", "", value))
                set_run_font(run, size=8.5, bold=True, color=RGBColor(255, 255, 255))
            else:
                if r_idx % 2 == 0:
                    shade_cell(cell, PALE)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) < 22 else WD_ALIGN_PARAGRAPH.LEFT
                add_inline(paragraph, value, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (("Heading 1", 15, 14, 7), ("Heading 2", 12.5, 11, 5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8.5)
        code.font.color.rgb = RGBColor(55, 65, 81)
        code.paragraph_format.left_indent = Inches(0.18)
        code.paragraph_format.right_indent = Inches(0.18)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(6)
        p_pr = code.element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F2F2")
        p_pr.append(shd)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(36)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Техническая спецификация набора данных для Power BI")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Перевозки и навигационные пломбы второго этапа")
    set_run_font(run, size=14, bold=False, color=RGBColor(55, 65, 81))

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = True
    values = [
        ("Проект", "ReportScriptEXCEL IKTT 2 STAGE"),
        ("Версия спецификации", "2.0"),
        ("Дата", "07.09.2026"),
        ("Начало второго этапа", "31.07.2026 00:00 UTC+05:00"),
    ]
    for i, (label, value) in enumerate(values):
        for j, item in enumerate((label, value)):
            cell = meta.cell(i, j)
            set_cell_margins(cell, 120, 130, 120, 130)
            set_cell_border(cell)
            if j == 0:
                shade_cell(cell, BLUE)
            run = cell.paragraphs[0].add_run(item)
            set_run_font(run, size=10, bold=(j == 0))

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Назначение документа")
    set_run_font(run, size=11, bold=True, color=RGBColor(0, 0, 0))
    p = doc.add_paragraph()
    add_inline(
        p,
        "Документ определяет входные данные, бизнес правила, выходные Excel таблицы, связи, коды и проверки, необходимые для создания модели Power BI. Его можно передать человеку или ИИ вместе с книгой технической спецификации и примерами выгрузок.",
        size=11,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run("Основное решение")
    set_run_font(run, size=11, bold=True, color=RGBColor(0, 0, 0))
    p = doc.add_paragraph()
    add_inline(
        p,
        "Power BI получает один стабильный набор таблиц из большой выгрузки, реестра пунктов и единого снимка пломб. Авто и ЖД определяются по большой выгрузке, а карта использует только реестровые и подтверждённые координаты.",
        size=11,
    )
    doc.add_page_break()


def parse_markdown(doc, lines):
    headings = [clean_heading(line[3:].strip()) for line in lines if line.startswith("## ")]
    doc.add_heading("Содержание", level=1)
    for number, item in enumerate(headings, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, f"{number}. " + re.sub(r"^\d+(?:\s+\d+)*\s+", "", item), size=10)
    doc.add_page_break()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if i == 0 and line.startswith("# "):
            i += 1
            continue
        if i < 7 and (not line or line.startswith("**Проект") or line.startswith("**Версия") or line.startswith("**Дата") or line.startswith("**Начало")):
            i += 1
            continue
        if line.startswith("## "):
            heading_text = clean_heading(line[3:].strip())
            # The final AI handoff is a self-contained instruction block.  Keep it
            # together on its own page instead of leaving its code block orphaned
            # after a mostly filled preceding page.
            if heading_text.startswith("20 "):
                doc.add_page_break()
            doc.add_heading(heading_text, level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(clean_heading(line[4:].strip()), level=2)
            i += 1
            continue
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            p = doc.add_paragraph(style="Code")
            p.add_run("\n".join(code_lines))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
            rows = []
            rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^[-*]\s+", "", line), size=10.25)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, line, size=10.25)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("```") or re.match(r"^[-*]\s+|^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        add_inline(p, " ".join(x.strip() for x in paragraph_lines), size=10.5)


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    parse_markdown(doc, lines)

    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.text = ""
        left = p.add_run("Power BI data specification  |  Версия 2.0  |  ")
        set_run_font(left, size=8, color=RGBColor(89, 89, 89))
        add_page_number(p)

    props = doc.core_properties
    props.title = "Техническая спецификация набора данных для Power BI"
    props.subject = "Перевозки и навигационные пломбы второго этапа"
    props.keywords = "Power BI, перевозки, пломбы, авто, железнодорожные перевозки"
    props.author = "Проект ReportScriptEXCEL IKTT"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
